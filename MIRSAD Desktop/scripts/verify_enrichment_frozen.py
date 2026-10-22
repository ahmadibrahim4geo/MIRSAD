"""Exercise optional climate outputs in the actual frozen worker, offline.

Uses the explicitly recorded public live fixture, not fabricated observations.
No installation, existing project deletion, or runtime replacement is performed.
"""
import argparse
import hashlib
import json
import os
import shutil
import subprocess
import sys
from pathlib import Path

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/"src"))


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--exe",type=Path,default=ROOT/".build/dist/MIRSAD/MIRSAD.exe")
    parser.add_argument("--cache",type=Path,default=ROOT/".test-output/climate-enrichment-quality-live/.mirsad-cache")
    parser.add_argument("--output",type=Path,default=ROOT/".test-output/frozen-enrichment-quality")
    args=parser.parse_args()
    exe=args.exe.resolve(strict=True);cache=args.cache.resolve(strict=True);output=args.output.resolve()
    if output.exists():raise FileExistsError("Preserve existing acceptance outputs; choose a new --output")
    from mirsad_desktop.daily_climate import VARIABLES,PM_PRODUCT
    from mirsad_desktop.climate_sources import ERA5_ONLY
    sources=[json.loads(p.read_text(encoding="utf-8")) for p in cache.glob("*.json") if not p.name.endswith("-summary.json")]
    if len(sources)!=1:raise ValueError("Expected the one-request public leap-month fixture cache")
    point=sources[0]["identity"]["station"]
    if point["latitude"]!=30 or point["longitude"]!=31 or point["usaf"]!="LIVE-0001":raise ValueError("Unexpected acceptance fixture")
    output.mkdir(parents=True)
    shutil.copytree(cache,output/".mirsad-cache")
    request=dict(operation="climate",points=[point],folder=str(output),name="Frozen_Source_Sample",metadata={},options=dict(
        start="2024-02-01",end="2024-02-29",provider=ERA5_ONLY,
        modules=["Temperature","Relative Humidity","Solar Radiation","Wind","Evapotranspiration"],
        extra_variables=list(VARIABLES),et_methods=["HG","PM"],elevation_options=dict(source="open-meteo")))
    job=output/"job";job.mkdir();path=job/"request.json"
    path.write_text(json.dumps(request,ensure_ascii=False),encoding="utf-8")
    env=dict(os.environ)
    for key in ("PYTHONPATH","PYTHONHOME"):env.pop(key,None)
    system=Path(env.get("SystemRoot","C:/Windows"))
    env.update(PATH=str(system/"System32")+";"+str(system),HTTPS_PROXY="http://127.0.0.1:9",HTTP_PROXY="http://127.0.0.1:9",NO_PROXY="")
    process=subprocess.run([str(exe),"--run-job",str(path)],cwd=Path.home(),env=env,
                           stdin=subprocess.DEVNULL,capture_output=True,timeout=120,creationflags=subprocess.CREATE_NO_WINDOW)
    (job/"worker.log").write_bytes(process.stdout+process.stderr)
    if process.returncode:raise RuntimeError("Frozen worker failed: "+str(job/"error.txt"))
    result=json.loads((job/"result.json").read_text(encoding="utf-8"))
    manifest=result["manifest"]
    if manifest["status"]!="complete":raise AssertionError(manifest["errors"])
    import pyogrio
    from openpyxl import load_workbook
    project=Path(result["output"])
    for filename in ("Project.gdb","Project.gpkg"):
        names=set(pyogrio.list_layers(project/filename)[:,0])
        assert {"Evapotranspiration",PM_PRODUCT}.issubset(names)
        meta,_,_,columns=pyogrio.raw.read(project/filename,layer="Relative_Humidity")
        row={name:column[0] for name,column in zip(meta["fields"],columns)}
        assert row["RHmax_Valid_Days"]==row["RHmin_Valid_Days"]==29 and row["Z_Source"]
    book=load_workbook(project/"15_Evapotranspiration/Penman_Monteith"/(PM_PRODUCT+".xlsx"),read_only=True,data_only=True)
    try:
        assert {"Daily","Month","Field_Dictionary","Methodology","Hourly_Processing","Elevation_Source"}.issubset(book.sheetnames)
        daily=list(book["Daily"].values)[1:]
        header,row=list(book["Month"].values)[:2]
        assert len(daily)==29 and abs(sum(values[-1] for values in daily)-dict(zip(header,row))["ET_PM_February_Mean"])<1e-8
        assert all(row[10:12]==(24,24) for row in list(book["Hourly_Processing"].values)[1:])
    finally:book.close()
    report=dict(passed=True,exe=str(exe),exe_sha256=hashlib.sha256(exe.read_bytes()).hexdigest(),
        scope="Actual frozen worker outside checkout, clean PATH, real recorded public source cache; network blocked by invalid proxy; no Toolbox/installer claim",
        checks=["Both method layers in real GPKG/FileGDB", "29 complete leap-month RH extremes and DEM source", "PM daily sums and hourly/DEM metadata in XLSX"],
        output=str(project))
    (output/"acceptance.json").write_text(json.dumps(report,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    print(json.dumps(report,ensure_ascii=True))


if __name__=="__main__":main()

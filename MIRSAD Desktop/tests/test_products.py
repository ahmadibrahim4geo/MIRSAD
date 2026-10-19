"""Native monthly product output/readback, quality and regional dispatch."""
import datetime as dt
import json
from pathlib import Path
import sqlite3
import sys
import tempfile
import unittest
from unittest.mock import patch
ROOT=Path(__file__).resolve().parents[1]
(ROOT/".test-output").mkdir(exist_ok=True)
sys.path.insert(0,str(ROOT/'src'))
from mirsad_desktop import climate,climate_products as p
from mirsad_desktop.inputs import make_point
from mirsad_desktop.services import run_climate_project
from mirsad_desktop.regional import run_regional_project


class ProductOutputs(unittest.TestCase):
    def test_native_monthly_zero_and_null_survive_excel_geopackage_filegdb(self):
        from openpyxl import load_workbook
        import pyogrio
        point=make_point('cell-A','Native grid cell',31.4,30.1)
        rec=dict(station=point,series={},monthly={'PRECTOTCORR':{(2024,1):0.}},
            native_monthly_coverage={'PRECTOTCORR':{'2024-01':1.}},status='OK',error='',
            resolution_by_variable={'PRECTOTCORR':{p.CHIRPS:dict(product='chirps3',spatial='0.05 degree',temporal='Monthly')}})
        with tempfile.TemporaryDirectory(dir=ROOT/'.test-output') as folder,patch.object(climate,'download',return_value=([rec],[],[])):
            output,manifest=run_climate_project([point],folder,'MonthlyContract',{},dt.date(2024,1,1),dt.date(2024,2,29),p.CHIRPS,['Precipitation'],save_daily=True)
            self.assertEqual(manifest['status'],'partial')
            book=load_workbook(output/'Precipitation.xlsx',read_only=True)
            self.assertNotIn('Daily',book.sheetnames);self.assertIn('Source_Resolution',book.sheetnames)
            rows=list(book['Data'].values);header,row=rows[:2]
            months=list(book['Month'].values);month_header,month_row=months[:2]
            book.close()
            for key,value in [('Coverage_Fraction',.5),('Valid_Days',None)]:self.assertEqual(row[header.index(key)],value)
            for key,value in [('R_January_Mean',0.),('R_February_Mean',None)]:self.assertEqual(month_row[month_header.index(key)],value)
            # Both spatial formats must contain the same calculated values/NULL.
            paths=list(output.rglob('*.gpkg'))+list(output.rglob('*.gdb'))
            self.assertEqual(len(paths),2,manifest['errors'])
            readback=0
            for path in paths:
                for layer,kind in pyogrio.list_layers(path):
                    info=pyogrio.read_info(path,layer=layer)
                    if 'R_January_Mean' not in info['fields']:continue
                    meta,_,geom,arrays=pyogrio.raw.read(path,layer=layer,columns=['R_January_Mean','R_February_Mean'])
                    self.assertEqual(float(arrays[0][0]),0.)
                    self.assertTrue(__import__('math').isnan(float(arrays[1][0])))
                    readback+=1
            self.assertEqual(readback,2,manifest['errors'])
            self.assertTrue(manifest['resolution_by_station'])

    def test_modis_primary_surface_coverage_does_not_require_optional_air(self):
        station=make_point('LST','Surface',31,30)
        rec=dict(station=station,series={},monthly={'TS':{(2024,1):26.85}},native_monthly_coverage={'TS':{'2024-01':2/31}},status='OK',error='')
        layers,daily=climate.prepare_outputs([rec],['Land Surface Temperature'],dt.date(2024,1,1),dt.date(2024,1,31),p.MODIS_DAY)
        props=layers['Land Surface Temperature'][0]['properties']
        self.assertEqual(props['LST_January_Mean'],26.85)
        self.assertIsNone(props['LST_minus_T_January'])
        self.assertAlmostEqual(props['Coverage_Fraction'],2/31)
        self.assertEqual(props['Status'],'PARTIAL');self.assertEqual(daily['Land Surface Temperature'][1],[])

    def test_regional_passes_original_source_cells_and_no_daily(self):
        from shapely.geometry import box,mapping
        cells=[dict(station_key='chirps3|2|3',name='native',longitude=10.025,latitude=20.025)]
        with tempfile.TemporaryDirectory(dir=ROOT/'.test-output') as folder,patch.object(p,'native_grid_cells',return_value=cells),patch('mirsad_desktop.services.run_climate_project') as run:
            output=Path(folder)/'Regional'
            def export(*args,**kwargs):
                output.mkdir();return output,{'processing':{}}
            run.side_effect=export
            run_regional_project(folder,'Regional',{},mapping(box(10,20,10.1,20.1)),None,'2024-01-01','2024-01-31',['Precipitation'],provider=p.CHIRPS)
            args,kwargs=run.call_args
            self.assertEqual(args[0][0]['longitude'],10.025)
            self.assertEqual(args[6],p.CHIRPS);self.assertFalse(kwargs['save_daily'])


if __name__=='__main__':unittest.main()

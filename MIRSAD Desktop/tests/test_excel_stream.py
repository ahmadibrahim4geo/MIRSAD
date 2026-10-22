"""Workbook row/part boundaries, metadata, literal strings and bounded row staging."""
import sys
import tempfile
import tracemalloc
import unittest
from pathlib import Path

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/"src"))
from mirsad_desktop.exports import workbook
from mirsad_desktop.excel_stream import ExcelWriteError
from mirsad_desktop.tabular import RowBuffer


class StreamingExcelTests(unittest.TestCase):
    def setUp(self):
        from openpyxl import load_workbook
        self.load=load_workbook
        (ROOT/".test-output").mkdir(exist_ok=True)
        self.tmp=tempfile.TemporaryDirectory(dir=ROOT/".test-output");self.addCleanup(self.tmp.cleanup)
        self.path=Path(self.tmp.name)/"Climate.xlsx"

    def test_sheet_and_workbook_parts_keep_every_row_once_with_metadata_and_index(self):
        data=[["A","2024-01-%02d"%(i+1),i] for i in range(11)]
        tables=[("Data",["Name"],[["العينة"]]),("Daily",["Point","Date","Value"],iter(data)),
                ("Source_Data",["Source"],iter([["Fixture"]])),("Methodology",["Rule"],[["UTC daily"]])]
        result=workbook(self.path,tables,sheet_capacity=3,part_daily_rows=4)
        self.assertTrue(result["split"])
        self.assertEqual(len(result["files"]),3)
        read=[]
        for path in result["files"]:
            book=self.load(path,read_only=True,data_only=False)
            try:
                self.assertEqual(book["Source_Data"].cell(2,1).value,"Fixture")
                self.assertIn("Methodology",book.sheetnames)
                for name in book.sheetnames:
                    if name=="Daily" or name.startswith("Daily_"):
                        self.assertLessEqual(book[name].max_row,4)
                        read.extend([list(row) for row in list(book[name].values)[1:]])
                if path==self.path:
                    self.assertIn("Export_Index",book.sheetnames)
                    rows=list(book["Export_Index"].values)[1:]
                    self.assertEqual(sum(row[2] for row in rows),11)
            finally:book.close()
        self.assertEqual(read,data)
        self.assertFalse(list(self.path.parent.glob("*.csv")))

    def test_small_export_keeps_sheet_order_dimensions_and_formula_strings_literal(self):
        workbook(self.path,[("Data",["Name","Value"],[["=1+1",None]]),("Daily",["Date"],[["2024-02-29"]]),("Source_Data",["Source"],[["Fixture"]])])
        book=self.load(self.path,read_only=True,data_only=False)
        try:
            self.assertEqual(book.sheetnames,["Data","Daily","Source_Data"])
            self.assertEqual(book["Data"].max_row,2)
            self.assertEqual(book["Data"].max_column,2)
            self.assertEqual(book["Data"].cell(2,1).value,"=1+1")
            self.assertEqual(book["Data"].cell(2,1).data_type,"s")
            self.assertIsNone(book["Data"].cell(2,2).value)
        finally:book.close()

    def test_successful_part_is_preserved_if_later_rows_fail(self):
        def broken():
            for i in range(9):yield [i]
            raise OSError("fixture source interrupted")
        with self.assertRaises(ExcelWriteError) as caught:
            workbook(self.path,[("Daily",["ID"],broken())],sheet_capacity=3,part_daily_rows=4)
        self.assertTrue(caught.exception.files)
        book=self.load(caught.exception.files[0],read_only=True,data_only=True)
        try:
            values=[row[0] for name in book.sheetnames for row in list(book[name].values)[1:]]
            self.assertEqual(values,[4,5,6,7])
        finally:book.close()

    def test_row_staging_rolls_to_disk_and_supports_repeated_readers(self):
        tracemalloc.start()
        rows=RowBuffer(memory_bytes=8192)
        try:
            for i in range(50000):rows.append([i,"x"*64])
            _,peak=tracemalloc.get_traced_memory()
            self.assertLess(peak,2*1024*1024)
            self.assertEqual(len(rows),50000)
            first,second=iter(rows),iter(rows)
            self.assertEqual(next(first),next(second))
            self.assertEqual(next(first),[1,"x"*64])
            self.assertEqual(next(second),[1,"x"*64])
            self.assertEqual(rows[-1],[49999,"x"*64])
        finally:rows.close();tracemalloc.stop()

    def test_excel_limits_are_validated_and_existing_workbooks_preserved(self):
        with self.assertRaises(ValueError):workbook(self.path,[("Data",["ID"],[])],sheet_capacity=1048576)
        self.path.write_bytes(b"original workbook bytes")
        with self.assertRaises(FileExistsError):workbook(self.path,[("Data",["ID"],[])])
        self.assertEqual(self.path.read_bytes(),b"original workbook bytes")


if __name__=="__main__":unittest.main()

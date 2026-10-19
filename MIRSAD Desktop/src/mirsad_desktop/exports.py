"""Real FileGDB/GeoPackage and Excel exports from a single set of records."""
from .crs import resolve_crs
import datetime as dt
import hashlib
import json
import math
import re
import uuid
import os
from pathlib import Path


def safe_name(value):
    return re.sub(r"[^A-Za-z0-9_-]+", "_", value).strip("_") or "MIRSAD_Project"


def station_properties(row):
    return {
        "Station_Key": row["station_key"], "Station_ID": row["usaf"], "WBAN": row["wban"],
        "Station_Name": row["name"], "Country": row["country"],
        "Source_Country": row.get("source_country", ""),
        "Point_Lat": row["latitude"], "Point_Lon": row["longitude"],
        "Elevation_m": row.get("elevation"), "Coordinate_QA": row.get("qa_status", "unverified")
        ,"Point_Type":row.get("point_type","Catalogue station"),"Input_Source":row.get("input_source","Embedded NOAA location catalogue")
    }


def write_layer(path, driver, layer, records, crs="EPSG:4326"):
    import numpy as np
    import pyogrio
    from pyogrio.raw import write
    from shapely import Point, to_wkb
    if "w" not in pyogrio.list_drivers().get(driver, ""):
        raise RuntimeError("Required writable driver unavailable: " + driver)
    from pyproj import CRS,Transformer
    layer=safe_name(layer)
    projection=resolve_crs(crs)
    transformer=Transformer.from_crs(4326,projection,always_xy=True)
    coordinates=[transformer.transform(r["longitude"],r["latitude"]) for r in records]
    if not all(math.isfinite(x) and math.isfinite(y) for x,y in coordinates):raise ValueError("Coordinates cannot be projected to "+str(crs))
    properties = [record["properties"] for record in records]
    fields = list(properties[0])
    arrays = []
    for field in fields:
        values = [row.get(field) for row in properties]
        valid = [v for v in values if v is not None]
        numeric = field not in ("Station_Key", "Station_ID", "WBAN", "Station_Name", "Country", "Source_Country", "Coordinate_QA",
                                "Data_Start", "Data_End", "Source", "Measurement_Unit", "Status", "Error_Msg","Point_Type","Input_Source","Gap_Fill_Method","Source_Temporal_Resolution")
        if numeric or (valid and all(isinstance(v, (int, float)) for v in valid)):
            arrays.append(np.array([float(v) if v is not None else np.nan for v in values], dtype="float64"))
        else:
            arrays.append(np.array(values, dtype=object))
    geometries = np.array([to_wkb(Point(x,y)) for x,y in coordinates], dtype=object)
    options = {"TARGET_ARCGIS_VERSION": "ALL"} if driver == "OpenFileGDB" else {}
    write(str(path), geometries, arrays, fields, driver=driver, layer=layer,
          geometry_type="Point", crs=projection.to_wkt(), encoding="UTF-8", layer_options=options)
    # Read back actual coordinates, stable keys, CRS and count after closing the writer.
    meta, fids, geometry, attributes = pyogrio.raw.read(str(path), layer=layer)
    from shapely import from_wkb
    if len(geometry) != len(records) or not meta.get("crs"):
        raise RuntimeError("Export readback failed: count or CRS mismatch")
    if not CRS.from_user_input(meta["crs"]).equals(projection,ignore_axis_order=True):raise RuntimeError("Export CRS changed during readback")
    key_index = list(meta["fields"]).index("Station_Key")
    expected = {r["properties"]["Station_Key"]: r for r in records}
    seen = set()
    for index, (key, geom) in enumerate(zip(attributes[key_index], geometry)):
        point, original = from_wkb(geom), expected[key]
        x,y=transformer.transform(original["longitude"],original["latitude"])
        tolerance=1e-4 if projection.is_projected else 1e-6
        if abs(point.x - x)>tolerance or abs(point.y-y)>tolerance:
            raise RuntimeError("Export coordinate mismatch: " + key)
        seen.add(key)
        for field_index, field in enumerate(meta["fields"]):
            actual, wanted = attributes[field_index][index], original["properties"].get(field)
            if wanted is None:
                if actual is not None and not (isinstance(actual, (float, np.floating)) and math.isnan(actual)):
                    raise RuntimeError("NULL changed during export: " + field)
            elif isinstance(wanted, (int,float)):
                if not math.isclose(float(actual),float(wanted),rel_tol=1e-9,abs_tol=1e-6):
                    raise RuntimeError("Numeric attribute changed during export: " + field)
            elif str(actual) != str(wanted):
                raise RuntimeError("Text attribute changed during export: " + field)
    if seen != set(expected):
        raise RuntimeError("Export station identity mismatch")
    return {"layer": layer, "features": len(records), "crs": meta["crs"], "readback": "passed"}


def element_sheets(records, source, catalogue_metadata, run_id, daily=None,source_resolutions=None):
    from .climate_products import descriptor
    product=descriptor(source)
    props = [row["properties"] for row in records]
    headers = list(props[0])
    monthly = [key for key in headers if re.search(r"_(January|February|March|April|May|June|July|August|September|October|November|December)_", key)]
    base = [key for key in headers if key not in monthly]
    sheets = [("Data",base,[[row.get(key) for key in base] for row in props])]
    if monthly:
        fields = ["Station_Key","Station_ID","Point_Lat","Point_Lon"] + monthly
        sheets.append(("Month",fields,[[row.get(key) for key in fields] for row in props]))
    sheets.append(("Source_Data",["Source","Catalogue","Run_ID","Start","End","Notes"],
        [[actual,catalogue_metadata.get("catalogue_version"),run_id,
          props[0].get("Data_Start"),props[0].get("Data_End"),
          "User input values; verify the original units." if source and source.startswith("User") else
          "Station locations only; no observations." if not source else
          "Selected product: "+product["id"]+"; native "+product["temporal"]+". Inspect Raw_Data for scaling, quality, availability and warnings. Monthly/composite values do not create daily observations. Unsupported dependencies remain NULL." if product else
          "Sources listed from actual values. Hybrid prioritizes NOAA GSOD (ISD-derived daily summaries) where station identity/location match, then selected ERA5 model, then NASA POWER. Raw responses, flags and per-value sources are retained in Raw_Data. NASA UV uses the legacy daylight-based noon-max approximation. Baseline: 1991–2020; trends need at least 8 complete years."]
          for actual in sorted({row.get("Source") or source or "Station catalogue (locations only)" for row in props})]))
    if daily:
        sheets.append(("Daily",daily[0],daily[1]))
    if source_resolutions:
        rows=[]
        for station_key,variables in source_resolutions.items():
            for variable,sources in variables.items():
                for actual,info in sources.items():rows.append([station_key,variable,actual,info.get("product"),info.get("spatial"),info.get("temporal")])
        sheets.append(("Source_Resolution",["Station_Key","Variable","Actual_Source","Product","Native_Resolution","Native_Temporal_Resolution"],rows))
    return sheets


def workbook(path, sheets):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    book = Workbook(write_only=False)
    book.remove(book.active)
    for name, headers, rows in sheets:
        # Keep complete daily records even beyond Excel's per-sheet capacity.
        capacity = 1000000
        for offset in range(0, max(1, len(rows)), capacity):
            suffix = "" if offset == 0 else "_" + str(offset // capacity + 1)
            ws = book.create_sheet(name[:31-len(suffix)] + suffix)
            ws.append(headers)
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="08273D")
                cell.alignment = Alignment(wrap_text=True)
            for row in rows[offset:offset+capacity]:
                # Raw strings are values, never Excel formulas from external data.
                ws.append(row)
                for cell in ws[ws.max_row]:
                    if cell.data_type == "f":
                        cell.data_type = "s"
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for column in ws.columns:
                ws.column_dimensions[column[0].column_letter].width = min(48, max(16, len(str(column[0].value)) + 3))
    book.save(path)


def export_project(stations, folder, name, catalogue_metadata, climate_layers=None, daily=None, source=None,
                   progress=None, cancelled=None, crs="EPSG:4326",source_resolutions=None):
    if not stations:
        raise ValueError("Choose at least one station.")
    progress = progress or (lambda message: None)
    cancelled = cancelled or (lambda: False)
    parent = Path(folder).resolve()
    parent.mkdir(parents=True, exist_ok=True)
    output = parent / safe_name(name)
    if output.exists():
        raise FileExistsError("Project already exists. Choose a new project name: " + str(output))
    output.mkdir()
    run_id = str(uuid.uuid4())
    from . import __version__
    manifest = dict(run_id=run_id,job_id=os.environ.get("MIRSAD_JOB_ID"), version=__version__,output_crs=crs, created_utc=dt.datetime.now(dt.timezone.utc).isoformat(),
                    catalogue=catalogue_metadata, outputs={}, errors=[], status="running")
    def save():
        (output / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    save()
    try:
        layers = {"Stations": [{"latitude": s["latitude"], "longitude": s["longitude"],
                                  "properties": station_properties(s)} for s in stations]}
        layers.update(climate_layers or {})
        total=len(layers)*3; completed=0
        for driver, filename in (("GPKG", "Project.gpkg"), ("OpenFileGDB", "Project.gdb")):
            try:
                for layer, records in layers.items():
                    if cancelled():
                        raise InterruptedError("Export cancelled")
                    progress({"stage":"export","completed":completed,"total":total,"message":"Exporting " + filename + " / " + layer})
                    result = write_layer(output / filename, driver, layer, records,crs=crs)
                    completed+=1
                    manifest["outputs"].setdefault(filename, []).append(result)
            except InterruptedError:
                raise
            except Exception as exc:
                manifest["errors"].append({"output": filename, "error": str(exc)})
            save()
        for layer, records in layers.items():
            if cancelled(): raise InterruptedError("Export cancelled")
            progress({"stage":"export","completed":completed,"total":total,"message":"Excel / " + layer})
            sheets = element_sheets(records,source,catalogue_metadata,run_id,(daily or {}).get(layer),source_resolutions=source_resolutions)
            path = output / (layer + ".xlsx")
            try:
                workbook(path, sheets)
                manifest["outputs"][path.name] = {"sha256": hashlib.sha256(path.read_bytes()).hexdigest()}
            except Exception as exc:
                manifest["errors"].append({"output":path.name,"error":str(exc)})
            completed+=1
            save()
        manifest["status"] = "partial" if manifest["errors"] else "complete"
        progress({"stage":"export","completed":total,"total":total,"message":"Exports complete"})
    except Exception as exc:
        manifest["status"] = "cancelled" if isinstance(exc, InterruptedError) else "failed"
        manifest["errors"].append({"error": str(exc)})
        raise
    finally:
        save()
    return output, manifest

"""Explicit small live acceptance: source -> optional DEM -> independent ET/Excel/GIS.

This fixture is not a default application country/site. No accounts/credentials.
Source acceptance is separate from frozen EXE and native Toolbox acceptance.
"""
import argparse
import datetime as dt
import json
import sys
from pathlib import Path

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/"src"))


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--live",action="store_true")
    parser.add_argument("--output",type=Path,default=ROOT/".test-output/climate-enrichment-live")
    args=parser.parse_args()
    if not args.live:parser.error("Network acceptance requires --live")
    from mirsad_desktop.inputs import make_point
    from mirsad_desktop.services import run_climate_project
    from mirsad_desktop.daily_climate import VARIABLES,PM_PRODUCT
    from mirsad_desktop.climate_sources import ERA5_ONLY
    import pyogrio
    from openpyxl import load_workbook
    args.output.mkdir(parents=True,exist_ok=True)
    point=make_point("LIVE-0001","Explicit historical source fixture",31,30)
    selected=list(VARIABLES)
    output,manifest=run_climate_project([point],str(args.output),"Verified_Source_Sample",{},
        dt.date(2024,2,1),dt.date(2024,2,29),ERA5_ONLY,
        ["Temperature","Relative Humidity","Solar Radiation","Wind","Evapotranspiration"],
        extra_variables=selected,et_methods=["HG","PM"],elevation_options=dict(source="open-meteo"),
        progress=lambda event:print(json.dumps(event,ensure_ascii=True),flush=True))
    checks=[]
    def check(value,message):
        if not value:raise AssertionError(message)
        checks.append(message)
    try:
        check(manifest["status"]=="complete","Complete source/service run without hidden export failures")
        for file in ("Project.gpkg","Project.gdb"):
            names=set(pyogrio.list_layers(output/file)[:,0])
            check({"Evapotranspiration",PM_PRODUCT}.issubset(names),file+": separate HG and PM layers")
            meta,_,_,columns=pyogrio.raw.read(output/file,layer="Relative_Humidity")
            row={name:column[0] for name,column in zip(meta["fields"],columns)}
            check(row["RHmax_Valid_Days"]==29 and row["RHmin_Valid_Days"]==29,file+": full leap-month hourly extrema coverage")
            check(row["Z_Source"] is not None and row["Z_Elev_m"] is not None,file+": actual DEM value/source retained")
            check("VPD_Source_February_Mean" in row and "VPD_FAO56_February_Mean" in row,file+": separate VPD statistics")
        path=output/"15_Evapotranspiration/Penman_Monteith"/(PM_PRODUCT+".xlsx")
        book=load_workbook(path,read_only=True,data_only=True)
        try:
            check({"Data","Month","Daily","Methodology","Field_Dictionary"}.issubset(book.sheetnames),"PM workbook has daily/monthly/method/field sheets")
            check({"Hourly_Processing","Elevation_Source","Elevation_Fields"}.issubset(book.sheetnames),"PM workbook includes hourly counts and DEM metadata")
            processing=list(book["Hourly_Processing"].values)[1:]
            check(bool(processing) and all(row[10:12]==(24,24) for row in processing),"Hourly processing exports complete 24-hour sample counts")
            elevation_header,elevation_row=list(book["Elevation_Source"].values)[:2]
            elevation=dict(zip(elevation_header,elevation_row))
            check(elevation["Surface_Type"]=="DSM" and "EGM2008" in elevation["Vertical_Datum"],"DEM workbook discloses DSM and vertical datum")
            daily=list(book["Daily"].values)[1:]
            header,row=list(book["Month"].values)[:2];month=dict(zip(header,row))
            check(len(daily)==29 and abs(sum(r[-1] for r in daily)-month["ET_PM_February_Mean"])<1e-8,
                  "Monthly PM is the sum of the 29 calculated daily values")
        finally:book.close()
        check(not any(name=="arcpy" or name.startswith("arcpy.") for name in sys.modules),"Desktop source runtime imports no ArcPy")
        report=dict(passed=True,scope="Small live Desktop-source acceptance; no frozen/native Toolbox claim",checks=checks,
                    output=str(output),source_warnings=manifest.get("source_warnings",[]),date="2026-10-21")
    except Exception as exc:
        report=dict(passed=False,checks=checks,error=str(exc),output=str(output),manifest_errors=manifest["errors"])
    (args.output/"acceptance.json").write_text(json.dumps(report,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    print(json.dumps(report,ensure_ascii=True))
    return 0 if report["passed"] else 1


if __name__=="__main__":raise SystemExit(main())

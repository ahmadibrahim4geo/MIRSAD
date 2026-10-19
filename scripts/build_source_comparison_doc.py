"""Rebuild the obsolete comparison document from the reviewed source snapshot.

Run with the Codex bundled Python for document authoring, then render and inspect.
"""
import json
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT=Path(__file__).resolve().parents[1]
catalogue=json.loads((ROOT/'data/climate_sources.json').read_text(encoding='utf-8'))
doc=Document();section=doc.sections[0]
section.page_width=Inches(8.5);section.page_height=Inches(11)
section.top_margin=section.bottom_margin=Inches(.7)
section.left_margin=section.right_margin=Inches(.8)
for name,size in [('Normal',10.5),('Title',23),('Heading 1',16),('Heading 2',13)]:
    style=doc.styles[name];style.font.name='Arial';style.font.size=Pt(size);style.font.color.rgb=RGBColor.from_string('151515')
    style.paragraph_format.space_after=Pt(6)
    fonts=style.element.get_or_add_rPr().find(qn('w:rFonts'))
    if fonts is not None:fonts.set(qn('w:cs'),'Arial')
for border in list(doc.styles.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
doc.styles['Normal'].paragraph_format.line_spacing=1.1
doc.core_properties.author='Ahmad Ibrahim'
doc.core_properties.title='MIRSAD climate data sources comparison'


def para(text,rtl=False,style=None):
    p=doc.add_paragraph(text,style=style)
    if rtl:
        p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
        bidi=OxmlElement('w:bidi');p._p.get_or_add_pPr().append(bidi)
        p._p.get_or_add_pPr().find(qn('w:jc')).set(qn('w:val'),'start')
        for run in p.runs:
            fonts=run._r.get_or_add_rPr();flag=OxmlElement('w:rtl');fonts.append(flag)
    return p


para('MIRSAD climate data sources comparison',style='Title')
para('مقارنة مصادر البيانات المناخية في مرصاد',rtl=True,style='Heading 1')
para('Ahmad Ibrahim  |  18 October 2026  |  Development 0.4.0.dev7')
para('This guide describes the implemented source adapters, native spatial and temporal resolution, archive coverage and account requirements shared by the ArcGIS toolboxes and MIRSAD Desktop. It replaces the previous five-source comparison, whose completeness and no-login CDS claims were inaccurate. Available variables and actual per-run provenance determine which indicators can be calculated.')
para('يوضح هذا الدليل المصادر المنفذة والدقة الأصلية والفترات الزمنية ومتطلبات الحساب المشتركة بين التول بوكس ومرصاد ديسكتوب. يستبدل المقارنة القديمة وادعاءات اكتمال البيانات والتنزيل من CDS دون دخول. المتغيرات المتاحة وبيانات التشغيل الفعلية تحدد المؤشرات الممكن حسابها.',rtl=True)
para('Automatic product selection',style='Heading 2')
para('Auto chooses the finest suitable product per variable inside the selected provider for the entire period. Open-Meteo uses ERA5-Land at 0.1 degree for supported temperature, humidity and dewpoint variables, and ERA5 at 0.25 degree for other supported variables. Intervals beginning before 1950 use ERA5 throughout. Advanced settings permit a fixed manual model. Fixed-grid products keep their native grids; smaller output pixels do not create source detail.')
para('يختار Auto أدق منتج مناسب لكل متغير داخل المصدر المختار وللفترة كاملة، مع اختيار يدوي متقدم عند الحاجة. شبكة المصدر الثابتة لا تتغير بتصغير خلية الراستر الناتج. يُسجل المنتج والدقة لكل متغير في بيانات المصدر الخام وورقة Source_Resolution في Desktop.',rtl=True)
para('Hybrid and missing data',style='Heading 2')
para('Station Hybrid finishes NOAA across the network, then ERA5 for missing values, then NASA POWER for remaining gaps, then optional monthly Gap Fill. Any-point Hybrid starts at ERA5. Raw observations and estimates remain separate. Missing precipitation is never zero; monthly/composite sources never fabricate daily observations. No source guarantees complete coverage.')
para('Need Sign In',style='Heading 2')
para('CDS needs the user’s personal API token and accepted dataset terms. IMERG needs NASA Earthdata credentials and GES DISC authorization. Credentials stay in memory and pass to Desktop workers through a pipe, outside project/job/log files. These adapters pass protocol and conversion checks; live authenticated end-to-end verification remains pending.')
para('يستخدم كل مستخدم حسابه الخاص؛ لا يحتاج البرنامج حساب المطور. تظل الأسرار خارج ملفات المشروع والسجلات. نهاية الأرشيف متغيرة، وتاريخ بداية المنتج لا يضمن وجود كل عنصر في كل موقع.',rtl=True)

sources=catalogue['sources']
for index,item in enumerate(sources):
    if index%2==0:doc.add_page_break()
    para(item['label'],style='Heading 1')
    para('Coverage  '+item['start']+' to '+item['end'])
    para('Native resolution  '+item['resolution']+'  |  '+item['temporal'])
    para('Region and tools  '+item['region']+'  |  '+', '.join(item['tools']))
    para('Variables  '+', '.join(item['variables']))
    para(item['notes_en'])
    para(item['notes_ar'],rtl=True)
    p=para('Official reference  '+item['url']);p.paragraph_format.space_after=Pt(16)
footer=section.footer.paragraphs[0]
footer.text='MIRSAD  |  Ahmad Ibrahim  |  Source catalogue '+catalogue['version']
footer.style=doc.styles['Normal'];footer.runs[0].font.size=Pt(8)
target=ROOT/'Documents/05_Climate_Data_Sources_Directory/Climate_Data_Sources_Comparison_Guide.docx'
doc.save(target)
print(target)

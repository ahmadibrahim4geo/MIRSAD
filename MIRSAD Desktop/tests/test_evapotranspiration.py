"""Separate ET products, accumulation semantics, optional rasters and missing Z."""
import datetime as dt
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/"src"))
from mirsad_desktop import climate,daily_climate
from mirsad_desktop.inputs import make_point
from mirsad_desktop.services import run_climate_project


class ETProductsTests(unittest.TestCase):
    def setUp(self):
        (ROOT/".test-output").mkdir(exist_ok=True)
        self.tmp=tempfile.TemporaryDirectory(dir=ROOT/".test-output");self.addCleanup(self.tmp.cleanup)
        self.start,self.end=dt.date(2024,1,1),dt.date(2024,12,31)
        self.point=make_point("0001","ET fixture",31,30,elevation=100)
        dates=[self.start+dt.timedelta(days=i) for i in range(366)]
        self.raw={p:{d:v for d in dates} for p,v in dict(T2M=16.9,T2M_MAX=21.5,T2M_MIN=12.3,
            RH2M_MAX=84,RH2M_MIN=63,WS10M_CONFIRMED=10/3.6,PS=100.1,ALLSKY_SFC_SW_DWN=22.07).items()}

    def run_project(self,methods,z=100):
        point=dict(self.point,elevation=z)
        record=dict(station=point,series=self.raw,status="OK",error="",source_summary="Fixture",provenance={})
        with patch.object(climate,"download",return_value=([record],[],[])), patch.object(daily_climate,"penman_monteith",return_value=4.):
            # Calculation is independently tested against FAO-56. This fixture
            # isolates export aggregation using an exact daily 4 mm reference.
            return run_climate_project([point],self.tmp.name,"ET_"+"_".join(methods),{},self.start,self.end,
                "NASA POWER",["Evapotranspiration"],et_methods=methods)

    def test_pm_only_is_a_separate_layer_and_workbook_with_daily_sums(self):
        import pyogrio
        from openpyxl import load_workbook
        output,manifest=self.run_project(["PM"])
        self.assertEqual(manifest["status"],"complete",manifest["errors"])
        names=set(pyogrio.list_layers(output/"Project.gpkg")[:,0])
        self.assertNotIn("Evapotranspiration",names)
        self.assertIn(daily_climate.PM_PRODUCT,names)
        meta,_,_,data=pyogrio.raw.read(output/"Project.gpkg",layer=daily_climate.PM_PRODUCT)
        row={k:v[0] for k,v in zip(meta["fields"],data)}
        self.assertEqual(row["ET_PM_Annual_Mean"],1464)
        self.assertEqual(row["ET_PM_February_Mean"],116)
        self.assertEqual(row["ET_PM_Winter_Mean"],364)
        path=output/"15_Evapotranspiration/Penman_Monteith"/(daily_climate.PM_PRODUCT+".xlsx")
        book=load_workbook(path,read_only=True,data_only=True)
        try:
            self.assertTrue({"Data","Month","Daily","Field_Dictionary","Methodology"}.issubset(book.sheetnames))
            self.assertTrue(all(r[4]=="ET0_PM" for r in list(book["Daily"].values)[1:]))
        finally:book.close()

    def test_hg_and_pm_keep_independent_fields_layers_and_directories(self):
        import pyogrio
        output,manifest=self.run_project(["HG","PM"])
        self.assertEqual(manifest["status"],"complete",manifest["errors"])
        for file in ("Project.gpkg","Project.gdb"):
            names=set(pyogrio.list_layers(output/file)[:,0])
            self.assertTrue({"Evapotranspiration",daily_climate.PM_PRODUCT}.issubset(names))
            hg=pyogrio.read_info(output/file,layer="Evapotranspiration")["fields"]
            pm=pyogrio.read_info(output/file,layer=daily_climate.PM_PRODUCT)["fields"]
            self.assertIn("ET_Annual_Mean",hg)
            self.assertNotIn("ET_PM_Annual_Mean",hg)
            self.assertIn("ET_PM_Annual_Mean",pm)
            self.assertNotIn("ET_Annual_Mean",pm)
        self.assertTrue((output/"15_Evapotranspiration/Hargreaves/Evapotranspiration.xlsx").is_file())
        self.assertTrue((output/"15_Evapotranspiration/Penman_Monteith"/(daily_climate.PM_PRODUCT+".xlsx")).is_file())

    def test_missing_known_height_is_not_silently_replaced_with_sea_level(self):
        from mirsad_desktop.enrichment import build_pm
        records=[dict(station=dict(self.point,elevation=None),series=self.raw,status="OK",error="",source_summary="fixture")]
        layer,daily,meta=build_pm(records,self.start,self.end,"fixture")
        self.assertEqual(layer[0]["properties"]["Status"],"NO_DATA")
        self.assertIsNone(layer[0]["properties"]["ET_PM_Annual_Mean"])
        self.assertEqual(daily[1],[])
        self.assertIsNone(records[0]["supplementary_provenance"]["ET0_PM"][0]["input_elevation_m"])

    def test_raster_selection_does_not_require_both_method_surfaces(self):
        from mirsad_desktop.analysis import build_atlas
        layers={name:[dict(latitude=30+y,longitude=31+x,properties={field:2.}) for x,y in ((0,0),(.1,0),(0,.1),(.1,.1))]
                for name,field in (("Evapotranspiration","ET_Annual_Mean"),(daily_climate.PM_PRODUCT,"ET_PM_Annual_Mean"))}
        output=Path(self.tmp.name);manifest=dict(outputs={},errors=[],run_id="ET-raster-fixture")
        build_atlas(layers,output,manifest,lambda e:None,lambda:False,crs="EPSG:32636",cell_size=2000,
                    fields=["ET_Annual_Mean","ET_PM_Annual_Mean"],pdf=False,et_raster_methods=["PM"])
        self.assertEqual(manifest["errors"],[])
        self.assertTrue(list((output/"Atlas/Penman_Monteith").glob("*.tif")))
        self.assertFalse((output/"Atlas/Hargreaves").exists())
        import rasterio
        with rasterio.open(next((output/"Atlas/Penman_Monteith").glob("*.tif"))) as raster:
            self.assertEqual(raster.tags()["MIRSAD_Unit"],"mm/year")

    def test_supplementary_raster_uses_field_units_instead_of_parent_percent(self):
        from mirsad_desktop.analysis import build_atlas
        import rasterio
        layers={"Relative Humidity":[dict(latitude=30+y,longitude=31+x,
            properties={"ea_Annual_Mean":1.7,"Measurement_Unit":"percent"})
            for x,y in ((0,0),(.1,0),(0,.1),(.1,.1))]}
        output=Path(self.tmp.name);manifest=dict(outputs={},errors=[],run_id="units-fixture")
        build_atlas(layers,output,manifest,lambda e:None,lambda:False,crs="EPSG:32636",cell_size=2000,pdf=False)
        self.assertEqual(manifest["errors"],[])
        with rasterio.open(next((output/"Atlas").glob("*.tif"))) as raster:
            self.assertEqual(raster.tags()["MIRSAD_Unit"],"kPa")

    def test_empty_unknown_and_duplicate_methods_are_rejected(self):
        for methods in ([],["unknown"],["PM","PM"]):
            with self.assertRaises(ValueError):daily_climate.validate_et_methods(methods,["Evapotranspiration"])

    def test_pm_output_units_and_bilingual_period_definitions(self):
        from mirsad_desktop.enrichment import build_pm
        results=[dict(station=self.point,series=self.raw,status="OK",error="",source_summary="Fixture")]
        with patch.object(daily_climate,"penman_monteith",return_value=25.4):
            records,daily,metadata=build_pm(results,self.start,self.end,"fixture",units="Imperial")
        self.assertAlmostEqual(records[0]["properties"]["ET_PM_Annual_Mean"],366)
        self.assertAlmostEqual(daily[1][0][6],1)
        self.assertEqual(daily[1][0][5],"inches/day")
        annual=next(r for r in metadata["dictionary"] if r[1]=="ET_PM_Annual_Mean")
        self.assertEqual(annual[4],"inches/year")
        self.assertEqual(annual[8],"Annual")
        self.assertIn("بنمان",annual[10])
        self.assertIn("G=0",annual[14])
        with patch.object(daily_climate,"penman_monteith",return_value=25.4):
            metric,_,_=build_pm(results,self.start,self.end,"fixture",units="Imperial",unit_options={"unit_precip":"Millimeters (mm)"})
        self.assertAlmostEqual(metric[0]["properties"]["ET_PM_February_Mean"],29*25.4)


if __name__=="__main__":unittest.main()

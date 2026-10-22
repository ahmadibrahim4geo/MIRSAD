"""Streaming XLSX tables and indexed Daily parts, without a CSV escape path."""
from pathlib import Path
import re
import shutil


class ExcelWriteError(RuntimeError):
    def __init__(self,message,files):
        super().__init__(message);self.files=list(files)


def write_workbook(path,sheets,sheet_capacity=1000000,part_daily_rows=4000000):
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment,Font,PatternFill
    from openpyxl.utils import get_column_letter
    path=Path(path)
    if path.suffix.lower()!=".xlsx":raise ValueError("Streaming writer requires an XLSX path")
    if not 1<=sheet_capacity<=1048575 or part_daily_rows<1:raise ValueError("Invalid Excel row limits")
    sheets=list(sheets)
    from .tabular import RowBuffer
    for i,(name,headers,rows) in enumerate(sheets):
        if name!="Daily" and iter(rows) is rows:
            buffered=RowBuffer()
            for row in rows:buffered.append(row)
            sheets[i]=(name,headers,buffered)
    if not sheets:raise ValueError("Workbook requires at least one table")
    if path.exists():raise FileExistsError("Existing workbook will be preserved: "+str(path))
    path.parent.mkdir(parents=True,exist_ok=True)
    index=[];files=[];dimensions={};main=None;current=None;current_path=path;part=1;part_rows=0;split=False
    header_font=Font(bold=True,color="FFFFFF");header_fill=PatternFill("solid",fgColor="08273D")

    def cell(ws,value,header=False):
        if isinstance(value,str) and len(value)>32767:raise ValueError("Excel text exceeds 32,767 characters")
        c=WriteOnlyCell(ws,value=value)
        if c.data_type=="f":c.data_type="s"  # untrusted strings are literal values
        if header:c.font=header_font;c.fill=header_fill;c.alignment=Alignment(wrap_text=True)
        return c

    def new_sheet(book,name,headers,position=None):
        if not headers or len(headers)>16384:raise ValueError("Invalid Excel column count")
        base=str(name)[:31];title=base;number=2;used={s.title.casefold() for s in book.worksheets}
        while title.casefold() in used:
            suffix="_"+str(number);title=base[:31-len(suffix)]+suffix;number+=1
        ws=book.create_sheet(title=title,index=position);ws.freeze_panes="A2"
        for i,h in enumerate(headers,1):ws.column_dimensions[get_column_letter(i)].width=min(48,max(16,len(str(h))+3))
        ws.append([cell(ws,h,True) for h in headers])
        row=dict(Workbook=None,Sheet=title,Records=0,First_Date=None,Last_Date=None)
        return ws,row

    def append(ws,headers,row,entry):
        row=list(row)
        if len(row)>len(headers):raise ValueError("Excel row has more values than its header")
        ws.append([cell(ws,v) for v in row]+[None]*(len(headers)-len(row)))
        entry["Records"]+=1
        if "Date" in headers:
            i=headers.index("Date")
            date=str(row[i]) if len(row)>i and row[i] is not None else None
            if date:
                entry["First_Date"]=min(date,entry["First_Date"]) if entry["First_Date"] else date
                entry["Last_Date"]=max(date,entry["Last_Date"]) if entry["Last_Date"] else date

    def finish(ws,headers,entry):
        ref="A1:%s%d"%(get_column_letter(len(headers)),entry["Records"]+1)
        ws.auto_filter.ref=ref;dimensions[id(ws)]=ref

    def prepare_book(filename):
        book=Workbook(write_only=True);daily_tables=[]
        # Prepare metadata before writing Daily bodies. Initial worksheet order
        # is unchanged; the Daily worksheet is created at its original position.
        for name,headers,rows in sheets:
            if name=="Daily":
                ws,entry=new_sheet(book,name,headers);entry["Workbook"]=filename.name;index.append(entry)
                daily_tables.append([headers,rows,ws,entry,1])
                continue
            ws,entry=new_sheet(book,name,headers);entry["Workbook"]=filename.name
            for row in rows:
                if entry["Records"]==sheet_capacity:
                    finish(ws,headers,entry);ws,entry=new_sheet(book,name,headers);entry["Workbook"]=filename.name
                append(ws,headers,row,entry)
            finish(ws,headers,entry)
        return book,daily_tables

    def save(book,filename):
        if filename.exists():raise FileExistsError("Existing workbook will be preserved: "+str(filename))
        temp=filename.with_name(filename.stem+".writing.xlsx")
        try:
            # openpyxl's write-only sheets omit <dimension>, which breaks
            # existing read-only clients' max_row/max_column contracts. Add the
            # known dimension to its closed temporary XML with bounded copying;
            # no worksheet cells are loaded back into RAM.
            for ws in book.worksheets:
                if not ws.closed:ws.close()
                xml=Path(ws._writer.out)
                stage=xml.with_name(xml.name+".dimensions")
                try:
                    with xml.open("rb") as src,stage.open("wb") as dst:
                        prefix=src.read(65536);dimension=('<dimension ref="'+dimensions[id(ws)]+'"/>').encode()
                        if b"<dimension " in prefix:prefix=re.sub(rb'<dimension\b[^>]*/>',dimension,prefix,count=1)
                        else:
                            position=prefix.find(b"<sheetViews")
                            if position<0:raise RuntimeError("Unexpected streaming worksheet header; cannot preserve dimensions")
                            prefix=prefix[:position]+dimension+prefix[position:]
                        dst.write(prefix);shutil.copyfileobj(src,dst,length=1024*1024)
                    with stage.open("rb") as src,xml.open("wb") as dst:
                        shutil.copyfileobj(src,dst,length=1024*1024)
                finally:
                    if stage.exists():stage.unlink()
            book.save(temp);temp.replace(filename);files.append(filename)
        finally:
            if temp.exists():temp.unlink()

    try:
        main,tables=prepare_book(path);current=main
        if len(tables)>1:raise ValueError("Use one logical Daily table per workbook")
        if tables:
            headers,rows,ws,entry,sheet_number=tables[0]
            for row in rows:
                if part_rows==part_daily_rows:
                    finish(ws,headers,entry)
                    if current is not main:save(current,current_path)
                    split=True;part+=1;part_rows=0
                    current_path=path.with_name(path.stem+"_Part_%03d.xlsx"%part)
                    current,new_tables=prepare_book(current_path)
                    headers,_,ws,entry,sheet_number=new_tables[0]
                elif entry["Records"]==sheet_capacity:
                    split=True;finish(ws,headers,entry);sheet_number+=1
                    position=current.worksheets.index(ws)+1
                    ws,entry=new_sheet(current,"Daily_"+str(sheet_number),headers,position)
                    entry["Workbook"]=current_path.name;index.append(entry)
                append(ws,headers,row,entry);part_rows+=1
            finish(ws,headers,entry)
            if current is not main:save(current,current_path)
        if split:
            headers=["Workbook","Sheet","Records","First_Date","Last_Date"]
            ws,entry=new_sheet(main,"Export_Index",headers)
            for record in index:append(ws,headers,[record[k] for k in headers],entry)
            finish(ws,headers,entry)
        save(main,path)
        return dict(files=[path]+[p for p in files if p!=path],split=split,index=index)
    except Exception as exc:
        # Successful parts survive a later failure and are returned to the
        # manifest through the exception; the caller reports incomplete export.
        for book in (current,main):
            if book:
                for ws in book.worksheets:
                    try:ws.close()
                    except Exception:pass
        raise ExcelWriteError(str(exc),files) from exc

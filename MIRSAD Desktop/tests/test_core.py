import copy
import datetime as dt
import json
import sqlite3
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/"src"))

from mirsad_desktop.catalogue import Catalogue
from mirsad_desktop.climate import summarize
from mirsad_desktop.exports import export_project
from mirsad_desktop.services import run_climate_project


class CoreTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(dir=ROOT/".test-output")
        self.folder = Path(self.tmp.name)
        self.catalogue = Catalogue()
        self.stations = self.catalogue.stations("EG")[:2]

    def tearDown(self): self.tmp.cleanup()

    def test_embedded_catalogue_matches_the_shared_toolbox_snapshot(self):
        import hashlib
        self.assertEqual(self.catalogue.digest(),hashlib.sha256((ROOT.parent/"data/stations.sqlite").read_bytes()).hexdigest())
        self.assertEqual(len(self.catalogue.stations("EG")),85)
        self.assertEqual(len(self.catalogue.stations("PS")),24)
        self.assertEqual(self.catalogue.metadata()["count"],28066)
        self.assertEqual(len({r["station_key"] for r in self.catalogue.stations("EG")}),85)
        for row in self.catalogue.stations("EG"):
            self.assertEqual(row["qa_status"],"reference_match")

    def test_annual_rainfall_respects_leap_year_and_missing_is_not_zero(self):
        start,end = dt.date(2024,1,1),dt.date(2024,12,31)
        series = {start+dt.timedelta(days=i):1 for i in range(366)}
        summary = summarize(series,start,end,"Precipitation")
        self.assertEqual(summary["R_Annual_Mean"],366)
        self.assertEqual(summary["R_February_Mean"],29)
        self.assertEqual(summary["Coverage_Fraction"],1)
        series[dt.date(2024,2,29)] = None
        summary = summarize(series,start,end,"Precipitation")
        self.assertIsNone(summary["R_Annual_Mean"])
        self.assertIsNone(summary["R_Period_Total"])
        series[dt.date(2024,2,29)] = 0
        self.assertEqual(summarize(series,start,end,"Precipitation")["R_Annual_Mean"],365)

    def test_partial_interval_is_explicit_and_never_becomes_a_complete_year(self):
        start,end = dt.date(2024,1,1),dt.date(2024,1,2)
        series = {start:10,end:20}
        values = summarize(series,start,end,"Temperature")
        self.assertEqual(values["T_Period_Mean"],15)
        self.assertIsNone(values["T_Annual_Mean"])
        self.assertIsNone(values["T_January_Mean"])

    def test_real_dual_exports_preserve_arabic_nulls_ids_coordinates_and_excel(self):
        from openpyxl import load_workbook
        stations = copy.deepcopy(self.stations)
        stations[0].update(name="محطة تجريبية",usaf="001234",elevation=None)
        output,manifest = export_project(stations,self.folder,"تجربة GIS",self.catalogue.metadata())
        self.assertEqual(manifest["status"],"complete",manifest)
        self.assertTrue((output/"Project.gdb").is_dir())
        self.assertTrue((output/"Project.gpkg").is_file())
        for path in ("Project.gdb","Project.gpkg"):
            self.assertEqual(manifest["outputs"][path][0]["features"],2)
        book = load_workbook(output/"Stations.xlsx",read_only=True)
        rows = list(book["Data"].values)
        self.assertIn("محطة تجريبية",rows[1])
        self.assertIn("001234",rows[1])
        book.close()
        self.assertFalse(list(output.glob("*.csv")))
        with self.assertRaises(FileExistsError):
            export_project(stations,self.folder,"تجربة GIS",self.catalogue.metadata())

    def test_failed_filegdb_export_is_partial_and_does_not_hide_successful_geopackage(self):
        from mirsad_desktop.exports import write_layer
        def fail_gdb(path,driver,*args,**kwargs):
            if driver == "OpenFileGDB": raise OSError("Simulated FileGDB lock")
            return write_layer(path,driver,*args,**kwargs)
        with patch("mirsad_desktop.exports.write_layer",side_effect=fail_gdb):
            output,manifest = export_project(self.stations,self.folder,"Partial",self.catalogue.metadata())
        self.assertEqual(manifest["status"],"partial")
        self.assertTrue((output/"Project.gpkg").is_file())
        self.assertEqual(manifest["errors"][0]["output"],"Project.gdb")

    def test_yearly_exports_recompute_and_keep_daily_month_and_source_sheets(self):
        from openpyxl import load_workbook
        station = self.stations[0]
        start,end = dt.date(2023,1,1),dt.date(2024,12,31)
        temp = {start+dt.timedelta(days=i):(10 if (start+dt.timedelta(days=i)).year==2023 else 30) for i in range(731)}
        rain = {day:1 for day in temp}
        fixture = [dict(station=station,series={"Temperature":temp,"Precipitation":rain},status="OK",error="")]
        original = copy.deepcopy(fixture)
        with patch("mirsad_desktop.climate.download",return_value=(fixture,[],["fixture://test"])) as download:
            output,manifest = run_climate_project([station],self.folder,"Climate",self.catalogue.metadata(),start,end,
                "NASA POWER",["Temperature","Precipitation"],separate_years=True)
        download.assert_called_once()
        self.assertEqual(manifest["status"],"complete",manifest["errors"])
        for year,value,days in ((2023,10,365),(2024,30,366)):
            book = load_workbook(output/"Temperature/Years"/("Temperature_%d.xlsx"%year),read_only=True)
            self.assertEqual(set(book.sheetnames),{"Data","Month","Daily","Source_Data","Source_Resolution"})
            rows = list(book["Data"].values)
            self.assertEqual(rows[1][rows[0].index("T_Annual_Mean")],value)
            self.assertEqual(book["Daily"].max_row,days+1)
            book.close()
        self.assertEqual(fixture,original)

    def test_ghazza_station_is_palestine_and_arab_league_aggregate(self):
        ps_stations = self.catalogue.stations("PS")
        ghazza = next((s for s in ps_stations if s["usaf"] == "623380"), None)
        self.assertIsNotNone(ghazza, "Station 623380 GHAZZA must be present under Palestine (PS)")
        self.assertEqual(ghazza["country"], "PS")
        self.assertEqual(ghazza["country_name"], "Palestine")
        self.assertEqual(ghazza["continent"], "Asia")
        self.assertTrue(ghazza["is_arab"])

        arab_stations = self.catalogue.stations("ARAB")
        self.assertEqual(len(arab_stations), 944)
        self.assertIn("623380", {s["usaf"] for s in arab_stations})

        arab_countries = self.catalogue.countries("Arab")
        self.assertEqual(arab_countries[0][0], "ARAB")
        self.assertEqual(arab_countries[0][3], 944)

    def test_hierarchical_crs_structure_and_validation(self):
        from mirsad_desktop.crs import build_crs_hierarchy, validate_and_describe_crs
        tree = build_crs_hierarchy()
        top_titles = [f["title"] for f in tree["folders"]]
        self.assertIn("Geographic Coordinate Systems", top_titles)
        self.assertIn("Projected Coordinate Systems", top_titles)

        # The complete tree now follows the installed Pro taxonomy, including
        # region/country/Zone Series folders, rather than a hand-written subset.
        def leaves(folder):
            return len(folder.get("items",[]))+sum(leaves(f) for f in folder.get("subfolders",[]))
        self.assertEqual(sum(leaves(f) for f in tree["folders"]),7519)

        # Validate Egypt Red Belt
        d_eg = validate_and_describe_crs("EPSG:22992")
        self.assertTrue(d_eg["valid"])
        self.assertTrue(d_eg["is_projected"])
        self.assertEqual(d_eg["projection_method"], "Transverse Mercator")
        self.assertEqual(len(d_eg["parameters"]), 5)

        # Validate Palestine Grid
        d_ps = validate_and_describe_crs("EPSG:28191")
        self.assertTrue(d_ps["valid"])
        self.assertTrue(d_ps["is_projected"])


if __name__ == "__main__": unittest.main()

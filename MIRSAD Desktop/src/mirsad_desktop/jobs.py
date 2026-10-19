"""Allowlisted JSON jobs. Native GIS processing runs outside the UI process."""
import datetime as dt
import faulthandler
import json
import math
import sys
import traceback
from pathlib import Path
from .validation import bounded_json,local_dataset,validate_points,MAX_POINTS

def clean(value):
    if isinstance(value,dict):return {str(k):clean(v) for k,v in value.items()}
    if isinstance(value,(list,tuple)):return [clean(v) for v in value]
    if hasattr(value,"item"):return clean(value.item())
    if isinstance(value,float) and not math.isfinite(value):return None
    if value is None or isinstance(value,(str,bool,int,float)):return value
    return str(value)


def read_credential_pipe():
    """PyInstaller windowed workers have no sys.stdin wrapper; use their pipe."""
    if sys.stdin:return sys.stdin.readline(32769)
    import ctypes
    from ctypes import wintypes
    kernel=ctypes.WinDLL("kernel32",use_last_error=True)
    kernel.GetStdHandle.argtypes=[wintypes.DWORD];kernel.GetStdHandle.restype=wintypes.HANDLE
    kernel.ReadFile.argtypes=[wintypes.HANDLE,wintypes.LPVOID,wintypes.DWORD,ctypes.POINTER(wintypes.DWORD),wintypes.LPVOID]
    kernel.ReadFile.restype=wintypes.BOOL
    handle=kernel.GetStdHandle(-10 & 0xffffffff);result=bytearray()
    while len(result)<=32768:
        buffer=ctypes.create_string_buffer(4096);count=wintypes.DWORD()
        ok=kernel.ReadFile(handle,buffer,4096,ctypes.byref(count),None)
        if not ok or count.value==0:break
        result.extend(buffer.raw[:count.value])
        if b"\n" in result:break
    return bytes(result).decode("utf-8")

def inspect_dataset(path,workdir,layer=None,source_crs=None):
    from .inputs import table_rows,vector_layers,vector_rows,read_boundary
    path=local_dataset(path)
    base=dict(path=str(path),layer=layer)
    if path.suffix.lower() in (".xlsx",".csv",".tsv"):
        headers,rows=table_rows(path)
        return dict(base,kind="point_mapping",headers=headers,rows=rows[:5],features=len(rows),crs="EPSG:4326",tabular=True)
    if path.suffix.lower() in (".tif",".tiff"):
        import numpy as np
        import rasterio
        from rasterio.vrt import WarpedVRT
        from PIL import Image
        from pyproj import Transformer
        with rasterio.open(path) as source:
            if source.driver!="GTiff":raise ValueError("Only genuine GeoTIFF raster inputs are accepted.")
            crs=source.crs or source_crs
            if not crs:return dict(base,kind="needs_crs",geometry="Raster")
            with WarpedVRT(source,src_crs=crs,crs="EPSG:3857") as view:
                width=min(1000,view.width);height=max(1,round(width*view.height/max(1,view.width)))
                if height>1000:width=max(1,round(width*1000/height));height=1000
                bands=[1,2,3] if view.count>=3 else [1]
                data=view.read(bands,out_shape=(len(bands),height,width),masked=True).astype("float64")
                valid=np.isfinite(data.filled(np.nan))&~np.ma.getmaskarray(data)
                channels=[]
                for band,mask in zip(data,valid):
                    values=band.data[mask]
                    low,high=np.percentile(values,[2,98]) if len(values) else (0,1)
                    normalized=np.nan_to_num((band.filled(0)-low)/max(high-low,1e-12))
                    channels.append((np.clip(normalized,0,1)*255).astype("uint8"))
                if len(channels)==1:
                    from matplotlib import colormaps
                    rgba=(colormaps["viridis"](channels[0]/255)*255).astype("uint8")
                else:rgba=np.dstack([*channels,np.ones((height,width),dtype="uint8")*255])
                rgba[:,:,3]=valid.all(axis=0)*220
                png=Path(workdir)/"raster_preview.png";Image.fromarray(rgba).save(png)
                transformer=Transformer.from_crs(3857,4326,always_xy=True)
                west,south=transformer.transform(view.bounds.left,view.bounds.bottom);east,north=transformer.transform(view.bounds.right,view.bounds.top)
        return dict(base,kind="raster",preview=str(png),bounds=[west,south,east,north],crs=str(crs),geometry="Raster")
    layers=vector_layers(path)
    if not layer:return dict(base,kind="layers",layers=layers)
    if layer not in [n for n,_ in layers]:raise ValueError("Selected layer is absent from the dataset.")
    import pyogrio
    info=pyogrio.read_info(path,layer=layer)
    geometry=str(info.get("geometry_type") or "Table")
    crs=info.get("crs") or source_crs
    base.update(crs=str(crs or ""),features=int(info.get("features",-1)),geometry=geometry)
    if geometry!="Table" and not crs:return dict(base,kind="needs_crs")
    if geometry=="Point":
        headers,rows,_=vector_rows(path,layer,source_crs)
        return dict(base,kind="point_mapping",headers=headers,rows=rows[:5],tabular=False)
    if "Polygon" in geometry:
        from shapely.geometry import mapping
        polygon=read_boundary(path,layer,source_crs)
        return dict(base,kind="boundary",boundary=mapping(polygon))
    if geometry=="Table":
        meta,_,_,data=pyogrio.raw.read(path,layer=layer,max_features=20,read_geometry=False)
        rows=[dict(zip(meta["fields"],values)) for values in zip(*data)]
        return dict(base,kind="table",headers=list(meta["fields"]),rows=rows)
    if "LineString" in geometry:
        from shapely import from_wkb,count_coordinates
        from shapely.geometry import mapping
        from shapely.ops import transform
        from pyproj import Transformer
        meta,_,geoms,_=pyogrio.raw.read(path,layer=layer,max_features=2001,columns=[])
        if len(geoms)>2000:raise ValueError("Line display is limited to 2,000 features; subset the layer first.")
        projection=Transformer.from_crs(crs,4326,always_xy=True)
        lines=[transform(projection.transform,from_wkb(g)) for g in geoms if g is not None]
        if sum(count_coordinates(g) for g in lines)>200000:raise ValueError("Line display exceeds vertex limit; simplify a copy first.")
        return dict(base,kind="vector",geometries=[mapping(g.simplify(.001,preserve_topology=True)) for g in lines])
    raise ValueError("Unsupported geometry type: "+geometry)


def preview_dataset(path,layer=None,sheet=None,offset=0):
    """Bounded, read-only pages. Formulas/macros and external links are never executed."""
    from .validation import MAX_FIELDS
    path=local_dataset(path)
    if isinstance(offset,bool) or not isinstance(offset,int) or not 0<=offset<=100000:
        raise ValueError("Preview offset must be between 0 and 100,000 rows.")
    base=dict(kind="preview_table",path=str(path),layer=layer,offset=offset)
    suffix=path.suffix.lower()
    if suffix==".xlsx":
        from openpyxl import load_workbook
        book=load_workbook(path,read_only=True,data_only=True,keep_links=False)
        try:
            names=book.sheetnames
            if len(names)>256:raise ValueError("Workbook exceeds the 256-sheet preview limit.")
            if sheet is not None and sheet not in names:raise ValueError("Workbook sheet not found.")
            ws=book[sheet or names[0]]
            if ws.max_column and ws.max_column>MAX_FIELDS:raise ValueError("Too many workbook columns to preview.")
            rows=iter(ws.iter_rows(min_row=1,max_row=1,values_only=True))
            headers=[str(v) if v is not None else "" for v in next(rows,())]
            values=list(ws.iter_rows(min_row=offset+2,max_row=offset+502,max_col=len(headers) or 1,values_only=True))
            # Declared worksheet dimensions can contain trailing blank formatted rows.
            while values and not any(v is not None for v in values[-1]):values.pop()
            return dict(base,sheets=names,sheet=ws.title,headers=headers,
                        rows=values[:500],more=len(values)>500)
        finally:book.close()
    if suffix in (".csv",".tsv"):
        import csv,itertools
        with path.open(encoding="utf-8-sig",newline="") as handle:
            sample=handle.read(8192);handle.seek(0)
            try:dialect=csv.Sniffer().sniff(sample,delimiters=",;\t")
            except csv.Error:dialect=csv.excel
            reader=csv.reader(handle,dialect=dialect)
            headers=next(reader,[])
            if not headers or len(headers)>MAX_FIELDS:raise ValueError("Invalid preview column count.")
            rows=list(itertools.islice(reader,offset,offset+501))
            rows=[(row+[None]*len(headers))[:len(headers)] for row in rows]
            return dict(base,headers=headers,rows=rows[:500],more=len(rows)>500)
    from .inputs import vector_layers
    layers=vector_layers(path)
    if layer is None:
        if not layers:raise ValueError("Dataset has no readable layers.")
        if len(layers)>1:return dict(base,kind="preview_layers",layers=[n for n,_ in layers])
        layer=layers[0][0]
    if layer not in [n for n,_ in layers]:raise ValueError("Dataset layer not found.")
    import pyogrio
    info=pyogrio.read_info(path,layer=layer)
    if len(info["fields"])>MAX_FIELDS:raise ValueError("Too many layer columns to preview.")
    meta,_,geometry,data=pyogrio.raw.read(path,layer=layer,skip_features=offset,max_features=501)
    rows=list(zip(*data)) if len(data) else [[] for _ in range(len(geometry) if geometry is not None else 0)]
    points=[]
    if geometry is not None and info.get("crs") and str(info.get("geometry_type"))=="Point":
        from shapely import from_wkb
        from pyproj import Transformer
        transform=Transformer.from_crs(info["crs"],4326,always_xy=True)
        for wkb in geometry[:500]:
            if wkb is not None:
                point=from_wkb(wkb)
                if point is not None and not point.is_empty:
                    x,y=transform.transform(point.x,point.y)
                    if math.isfinite(x) and math.isfinite(y) and -180<=x<=180 and -90<=y<=90:points.append([x,y])
    return dict(base,layer=layer,headers=list(meta["fields"]),rows=rows[:500],
                more=len(rows)>500,crs=info.get("crs"),points=points)

def execute(request,workdir,progress,cancelled,credentials=None):
    operation=request.get("operation")
    if operation=="pdf_preview":
        path=Path(request["options"]["path"]).resolve(strict=True)
        page=request["options"].get("page",0)
        if path.suffix.lower()!=".pdf" or not path.is_file() or path.stat().st_size>100*1024*1024:
            raise ValueError("PDF preview requires a local PDF smaller than 100 MB.")
        if isinstance(page,bool) or not isinstance(page,int) or not 0<=page<500:raise ValueError("Invalid PDF page.")
        from PySide6.QtGui import QGuiApplication
        from PySide6.QtCore import QSize
        from PySide6.QtPdf import QPdfDocument
        app=QGuiApplication.instance() or QGuiApplication(["MIRSAD PDF worker","-platform","offscreen"])
        document=QPdfDocument()
        try:
            if document.load(str(path))!=QPdfDocument.Error.None_:raise ValueError("PDF could not be opened.")
            if not 0<=page<document.pageCount() or document.pageCount()>500:raise ValueError("PDF preview supports 1 to 500 pages.")
            size=document.pagePointSize(page)
            ratio=1600/max(1,size.width(),size.height())
            image=document.render(page,QSize(max(1,round(size.width()*ratio)),max(1,round(size.height()*ratio))))
            preview=Path(workdir)/"pdf_page.png"
            if image.isNull() or not image.save(str(preview)):raise ValueError("PDF page rendering failed.")
            return dict(kind="pdf_preview",path=str(path),page=page,pages=document.pageCount(),preview=str(preview))
        finally:document.close()
    if operation=="preview":return preview_dataset(**request["options"])
    if operation=="inspect":return inspect_dataset(workdir=workdir,**request["options"])
    if operation=="import_points":
        from .inputs import import_points
        points=import_points(**request["options"])
        options=request["options"];path=local_dataset(options["path"])
        source_crs=options.get("source_crs") or "EPSG:4326"
        if path.suffix.lower() not in (".csv",".tsv",".xlsx"):
            import pyogrio
            source_crs=pyogrio.read_info(path,layer=options.get("layer")).get("crs") or options.get("source_crs")
        return {"kind":"points","points":points,"path":str(path),"layer":options.get("layer"),
                "crs":str(source_crs) if source_crs else None}
    if operation=="regional":
        from .regional import run_regional_project
        result=run_regional_project(request["folder"],request["name"],request.get("metadata",{}),progress=progress,cancelled=cancelled,**request["options"])
        return {"kind":"project","output":str(result[0]),"manifest":result[1]}
    if operation not in ("export","climate","analysis"):raise ValueError("Unsupported job operation.")
    points=validate_points(request["points"])
    options=dict(request.get("options",{}));metadata=request.get("metadata",{})
    from .exports import export_project
    if operation=="export":result=export_project(points,request["folder"],request["name"],metadata,progress=progress,cancelled=cancelled,**options)
    elif operation=="climate":
        from .services import run_climate_project
        from shapely.geometry import shape
        for key in ("start","end"):options[key]=dt.date.fromisoformat(options[key])
        if options.get("raster_options",{}):
            boundary=options["raster_options"].get("boundary")
            if boundary:options["raster_options"]["boundary"]=shape(boundary)
        result=run_climate_project(points,request["folder"],request["name"],metadata,progress=progress,cancelled=cancelled,credentials=credentials,**options)
    else:
        from .analysis import run_existing_project
        from shapely.geometry import shape
        if options.get("boundary"):options["boundary"]=shape(options["boundary"])
        result=run_existing_project(points,request["folder"],request["name"],metadata,progress=progress,cancelled=cancelled,**options)
    return {"kind":"project","output":str(result[0]),"manifest":result[1]}

def run_request(path):
    path=Path(path).resolve();workdir=path.parent
    def emit(value):
        if sys.stdout:print(json.dumps(clean(value),ensure_ascii=True),flush=True)
    def state(status):
        (workdir/"state.json").write_text(json.dumps({"status":status,"utc":dt.datetime.now(dt.timezone.utc).isoformat()}),encoding="utf-8")
    with (workdir/"native_crash.txt").open("w",encoding="utf-8") as crash_log:
        faulthandler.enable(file=crash_log)
        try:
            state("running");request=bounded_json(path)
            credentials=None
            if request.get("requires_credentials"):
                payload=read_credential_pipe()
                if len(payload)>32768:raise ValueError("Invalid credential input")
                credentials=json.loads(payload)
                if not isinstance(credentials,dict) or set(credentials)-{"username","password","token"}:raise ValueError("Invalid credential fields")
                if any(not isinstance(v,str) or len(v)>8192 for v in credentials.values()):raise ValueError("Invalid credential values")
            result=execute(request,workdir,lambda update:emit({"type":"progress","value":update}),lambda:(workdir/"cancel").exists(),credentials=credentials)
            (workdir/"result.json").write_text(json.dumps(clean(result),ensure_ascii=False,allow_nan=False),encoding="utf-8")
            state("finished");emit({"type":"result"});return 0
        except BaseException as exc:
            (workdir/"error.txt").write_text(traceback.format_exc(),encoding="utf-8")
            emit({"type":"error","message":str(exc)[:2000]});state("cancelled" if isinstance(exc,InterruptedError) else "failed");return 2
        finally:faulthandler.disable()

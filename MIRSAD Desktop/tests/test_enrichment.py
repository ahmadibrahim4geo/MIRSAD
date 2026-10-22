"""Real optional-variable exports, independent expected values and failure behavior."""
import copy
import datetime as dt
import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/"src"))
from mirsad_desktop.inputs import make_point
from mirsad_desktop import climate
from mirsad_desktop.services import run_climate_project
from mirsad_desktop.enrichment import summarize


def fixture(missing=False):
    point=make_point("0001","اختبار الرطوبة",31,30,elevation=100)
    start,end=dt.date(2024,1,1),dt.date(2024,12,31)
    dates=[start+dt.timedelta(days=i) for i in range(366)]
    raw={p:{d:v for d in dates} for p,v in dict(RH2M=68.,RH2M_MAX=82.,RH2M_MIN=54.,
        T2M_MAX=25.,T2M_MIN=18.,VPD_SOURCE=3.).items()}
    raw["T2MDEW"]={}
    if missing:raw["RH2M_MIN"].pop(dt.date(2024,2,29))
    record=dict(station=point,series=raw,status="OK",error="",source_summary="Open-Meteo era5_land",
                provenance={p:{str(d):"Open-Meteo era5_land" for d in values} for p,values in raw.items()},
                processing_by_variable={p:{str(d):dict(operation=operation,valid_samples=24,expected_samples=24,
                    timezone="UTC",source="Open-Meteo era5_land",input_temporal="hourly",output_temporal="daily") for d in raw[p]}
                    for p,operation in (("RH2M_MAX","max"),("RH2M_MIN","min"))})
    return point,[record],start,end


class SupplementaryOutputsTests(unittest.TestCase):
    def setUp(self):
        (ROOT/".test-output").mkdir(exist_ok=True)
        self.tmp=tempfile.TemporaryDirectory(dir=ROOT/".test-output")
        self.addCleanup(self.tmp.cleanup)

    def run_project(self,missing=False,daily=True):
        point,records,start,end=fixture(missing)
        original=copy.deepcopy(records)
        selected=["RHmax","RHmin","ea","es","VPD_Source","VPD_FAO56"]
        with patch.object(climate,"download",return_value=(records,[],[])) as download:
            output,manifest=run_climate_project([point],self.tmp.name,"Optional_RH",{},start,end,
                "Open-Meteo ERA5",["Relative Humidity"],save_daily=daily,extra_variables=selected)
        self.assertEqual(records,original)
        self.assertEqual(download.call_args.kwargs["extra_variables"],selected)
        return output,manifest

    def test_daily_humidity_and_two_vpds_survive_real_excel_and_dual_gis(self):
        import pyogrio
        from openpyxl import load_workbook
        output,manifest=self.run_project()
        self.assertEqual(manifest["status"],"complete",manifest["errors"])
        for path in (output/"Project.gpkg",output/"Project.gdb"):
            meta,_,_,data=pyogrio.raw.read(path,layer="Relative_Humidity")
            values={name:column[0] for name,column in zip(meta["fields"],data)}
            self.assertEqual(values["RHmax_Annual_Mean"],82)
            self.assertEqual(values["RHmin_Annual_Mean"],54)
            self.assertAlmostEqual(values["ea_Annual_Mean"],1.70,places=2)
            self.assertAlmostEqual(values["VPD_FAO56_Annual_Mean"],.91,places=2)
            self.assertEqual(values["VPD_Source_Annual_Mean"],3)
        book=load_workbook(output/"Relative Humidity.xlsx",read_only=True,data_only=True)
        try:
            self.assertTrue({"Data","Month","Daily","Source_Data","Field_Dictionary","Methodology"}.issubset(book.sheetnames))
            methods=list(book["Methodology"].values)
            self.assertTrue(any("max" in str(row[5]) for row in methods[1:]))
            for row in list(book["Field_Dictionary"].values)[1:]:self.assertEqual(row[1],row[2])
            hourly=list(book["Hourly_Processing"].values)
            self.assertTrue(all(row[10:12]==(24,24) for row in hourly[1:]))
            self.assertTrue(any("مشتقة" in row[16] for row in hourly[1:]))
            daily=list(book["Daily"].values)
            names={row[4] for row in daily[1:]}
            self.assertTrue({"RHmax","RHmin","ea","es","VPD_Source","VPD_FAO56"}.issubset(names))
        finally:book.close()

    def test_daily_saving_is_optional_without_disabling_the_calculation(self):
        from openpyxl import load_workbook
        output,manifest=self.run_project(daily=False)
        self.assertEqual(manifest["status"],"complete")
        book=load_workbook(output/"Relative Humidity.xlsx",read_only=True,data_only=True)
        try:
            self.assertNotIn("Daily",book.sheetnames)
            headers=next(book["Data"].values)
            self.assertIn("ea_Annual_Mean",headers)
        finally:book.close()

    def test_missing_leap_day_is_partial_and_not_zero_or_a_complete_annual_mean(self):
        import pyogrio
        output,manifest=self.run_project(missing=True)
        self.assertEqual(manifest["status"],"partial")
        self.assertTrue(any(e.get("variable")=="ea" for e in manifest["errors"]))
        meta,_,_,data=pyogrio.raw.read(output/"Project.gpkg",layer="Relative_Humidity")
        values={name:column[0] for name,column in zip(meta["fields"],data)}
        self.assertEqual(values["ea_Valid_Days"],365)
        import math
        self.assertTrue(math.isnan(values["ea_Annual_Mean"]))
        self.assertEqual(values["RHmax_Annual_Mean"],82)

    def test_parent_validation_and_native_monthly_rejection_happen_before_exports(self):
        point,_,start,end=fixture()
        with self.assertRaisesRegex(ValueError,"parent module"):
            climate.download([point],start,end,"NASA POWER",["Temperature"],self.tmp.name,
                             lambda e:None,lambda:False,extra_variables=["RHmax"])
        from mirsad_desktop.climate_products import CHELSA
        with self.assertRaisesRegex(ValueError,"not integrated"):
            climate.download([point],start,end,CHELSA,["Relative Humidity"],self.tmp.name,
                             lambda e:None,lambda:False,extra_variables=["RHmax"])

    def test_new_summary_names_distinguish_mean_daily_extremes_and_interannual_range(self):
        values={dt.date(2023,1,d):80. for d in range(1,32)}
        values.update({dt.date(2024,1,d):90. for d in range(1,32)})
        fields=summarize(values,dt.date(2023,1,1),dt.date(2024,12,31),"RHmax",ranges=True)
        self.assertEqual(fields["RHmax_January_Mean"],85)
        self.assertEqual(fields["RHmax_January_Range"],10)
        self.assertIsNone(fields["RHmax_Annual_Mean"])

    def test_invalid_hourly_dewpoint_uses_valid_rh_fallback_without_changing_raw(self):
        from mirsad_desktop.enrichment import calculation_inputs
        from mirsad_desktop.daily_climate import derive_day
        _,records,start,_=fixture()
        result=records[0];result["series"]["T2MDEW"][start]=0.
        result["processing_by_variable"]["T2MDEW"]={str(start):dict(status="invalid_hourly_input",expected_samples=24)}
        original=copy.deepcopy(result)
        values,metadata=derive_day(calculation_inputs(result,start),30,start,elevation=100)
        self.assertAlmostEqual(values["ea"],1.70,places=2)
        self.assertEqual(metadata["ea_method"],"FAO56_eq17_RH_extrema")
        self.assertEqual(result,original)


if __name__=="__main__":unittest.main()

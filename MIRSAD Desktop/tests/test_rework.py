"""Independent checks for the requested UI/data-contract repairs."""
import csv
import datetime as dt
import io
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

ROOT=Path(__file__).resolve().parents[1]
(ROOT/".test-output").mkdir(exist_ok=True)
sys.path.insert(0,str(ROOT/"src"))
from mirsad_desktop import climate
from mirsad_desktop.inputs import make_point
from mirsad_desktop.observations import parse_gsod


class ReworkTests(unittest.TestCase):
    def test_noaa_units_flags_identity_and_location(self):
        station=dict(usaf="012345",wban="00001",latitude=30,longitude=31)
        header=["STATION","DATE","LATITUDE","LONGITUDE","TEMP","MAX","MIN","PRCP","PRCP_ATTRIBUTES","WDSP","SLP","STP"]
        values=[["01234500001","2024-01-01",30,31,68,86,50,1,"G",10,1000,990],
                ["01234500001","2024-01-02",30,31,9999.9,9999.9,9999.9,0,"I",999.9,9999.9,9999.9],
                ["01234500001","2024-01-03",30,31,68,86,50,0,"H",10,1000,990],
                ["01234500001","2024-01-04",31,31,68,86,50,2,"G",10,1000,990]]
        text=io.StringIO();writer=csv.writer(text);writer.writerow(header);writer.writerows(values)
        result,rejected=parse_gsod(text.getvalue(),station,dt.date(2024,1,1),dt.date(2024,1,4),
                                  ["T2M","T2M_MAX","T2M_MIN","PRECTOTCORR","WS10M","SLP","PS"])
        day=dt.date(2024,1,1)
        self.assertEqual(result["T2M"][day],20)
        self.assertEqual(result["T2M_MAX"][day],30)
        self.assertEqual(result["T2M_MIN"][day],10)
        self.assertEqual(result["PRECTOTCORR"],{day:25.4})
        self.assertAlmostEqual(result["WS10M"][day],5.144444444444445)
        self.assertEqual(result["SLP"][day],100);self.assertEqual(result["PS"][day],99)
        self.assertEqual(rejected,1)
        with self.assertRaises(ValueError):parse_gsod(text.getvalue().replace("01234500001","01234500002"),station,day,day,["T2M"])

    def test_selected_openmeteo_model_reaches_actual_request(self):
        station=make_point("p","Point",31,30);day=dt.date(2024,1,1)
        def response(session,url,query,cancelled):
            self.assertEqual(query["models"],"era5_land")
            return {"daily":{"time":["2024-01-01"],"temperature_2m_mean":[20],"temperature_2m_max":[25],"temperature_2m_min":[15]}},url
        with patch.object(climate,"request_json",response):
            series,_,_=climate.fetch_series(None,station,day,day,"Open-Meteo ERA5",lambda:False,["Temperature"],"ERA5-Land (Highest Resolution)")
        self.assertEqual(series["T2M"][day],20)

    def test_hybrid_preserves_observations_and_fills_missing_dates(self):
        day=dt.date(2024,1,1);next_day=day+dt.timedelta(days=1)
        station=dict(usaf="012345",wban="00001",latitude=30,longitude=31)
        def response(url,session,cancelled):
            if "global-summary-of-the-day" in url:
                return "STATION,DATE,LATITUDE,LONGITUDE,PRCP,PRCP_ATTRIBUTES\n01234500001,2024-01-01,30,31,1,G\n"
            if "open-meteo" in url:
                return json.dumps({"daily":{"time":[str(day),str(next_day)],"precipitation_sum":[9,None]}})
            return json.dumps({"properties":{"parameter":{"PRECTOTCORR":{"20240101":8,"20240102":2}}}})
        with patch("mirsad_desktop.hybrid_pipeline._read",response):
            series,_,payload=climate.fetch_series(None,station,day,next_day,"Hybrid",lambda:False,["Precipitation"])
        self.assertEqual(series["PRECTOTCORR"],{day:25.4,next_day:2})
        self.assertIn("NOAA",payload["provenance"]["PRECTOTCORR"][str(day)])
        self.assertEqual(payload["provenance"]["PRECTOTCORR"][str(next_day)],"NASA POWER")
        self.assertEqual([r["source"] for r in payload["raw_responses"]],["NOAA GSOD","Open-Meteo","NASA POWER"])

    def test_season_selection_keeps_monthly_workbook_fields(self):
        day=dt.date(2024,1,1);point=make_point("p","Point",31,30)
        series={p:{day+dt.timedelta(days=i):20 for i in range(366)} for p in ("T2M","T2M_MAX","T2M_MIN")}
        layers,_=climate.prepare_outputs([dict(station=point,series=series,status="OK",error="")],["Temperature"],day,dt.date(2024,12,31),"NASA POWER",scopes=["Winter"])
        props=layers["Temperature"][0]["properties"]
        self.assertEqual(props["T_Winter_Mean"],20)
        self.assertEqual(props["T_January_Mean"],20)
        self.assertNotIn("T_Summer_Mean",props);self.assertNotIn("T_Annual_Mean",props)

    def test_official_crs_catalogue_has_definitions_and_correct_types(self):
        from mirsad_desktop.crs import crs_catalogue,resolve_crs,build_crs_hierarchy
        rows=crs_catalogue();self.assertGreater(len(rows),8000)
        self.assertEqual(len({r["code"] for r in rows}),len(rows))
        self.assertTrue(resolve_crs("EPSG:4087").is_projected)
        self.assertTrue(resolve_crs("EPSG:4326").is_geographic)
        tree=build_crs_hierarchy();self.assertEqual(len(tree["folders"]),2)

    def test_crs_native_folders_survive_additional_catalogue_and_keep_utm_depth(self):
        from mirsad_desktop.crs import build_crs_hierarchy
        geographic=["3D","Africa","Antarctica","Asia","Australia and New Zealand","Central America",
                    "Europe","North America","Oceans","Solar System","South America","World"]
        projected=["ARC (equal arc-second)","Continental","County Systems","Gauss Kruger","National Grids",
                   "Polar","Solar System","State Plane","State Systems","Tribal","UTM","World","World (Sphere-based)"]
        utm=["Africa","Asia","BLM (US Feet)","Europe","Indonesia","Italy","Malaysia","New Zealand",
             "North America","Oceans","South America","WGS 1972 BE","WGS 1972","WGS 1984"]
        for latest in (False,True):
            roots=build_crs_hierarchy(latest=latest)["folders"]
            self.assertEqual([r["title"] for r in roots],["Geographic Coordinate Systems","Projected Coordinate Systems"])
            self.assertEqual([r["title"] for r in roots[0]["subfolders"]][:12],geographic)
            self.assertEqual([r["title"] for r in roots[1]["subfolders"]][:13],projected)
            utm_folder=next(r for r in roots[1]["subfolders"] if r["title"]=="UTM")
            self.assertEqual([r["title"] for r in utm_folder["subfolders"]],utm)
            wgs=next(r for r in utm_folder["subfolders"] if r["title"]=="WGS 1984")
            self.assertEqual([r["title"] for r in wgs["subfolders"]],["Northern Hemisphere","Southern Hemisphere"])
            self.assertIn("EPSG:32636",[r[0] for r in wgs["subfolders"][0]["items"]])
            self.assertIn("EPSG:32736",[r[0] for r in wgs["subfolders"][1]["items"]])
            def leaves(folder):
                return list(folder.get("items",[]))+[row for sub in folder.get("subfolders",[]) for row in leaves(sub)]
            codes=[row[0].upper() for root in roots for row in leaves(root)]
            self.assertEqual(len(codes),9279 if latest else 7519)
            self.assertEqual(len(set(codes)),9185 if latest else 7425)

    def test_interpolation_expected_idw_and_reusable_index(self):
        import numpy as np
        from mirsad_desktop.analysis import Interpolator
        xy=np.array([[0,0],[2,0],[0,2]],float);values=np.array([0,20,40],float)
        predict=Interpolator(xy,values,power=2,neighbours=3)
        # Equal distance from all three samples gives their arithmetic mean.
        self.assertAlmostEqual(predict([[1,1]])[0],20)
        self.assertAlmostEqual(predict([[2,0]])[0],20)

    def test_regional_monthly_precipitation_and_month_13(self):
        from mirsad_desktop.regional import parse_features
        start,end=dt.date(2024,1,1),dt.date(2024,12,31)
        payload={"features":[{"geometry":{"coordinates":[31,30]},"properties":{"parameter":{"PRECTOTCORR":{"202401":1,"202402":2,"202403":-999,"202413":12}}}}]}
        lon,lat,monthly=next(parse_features(payload,"PRECTOTCORR",start,end))
        self.assertEqual(monthly,{(2024,1):31,(2024,2):58})
        point=make_point("grid","Grid",lon,lat)
        layers,daily=climate.prepare_outputs([dict(station=point,series={},monthly={"PRECTOTCORR":monthly},status="OK",error="")],["Precipitation"],start,end,"NASA POWER Regional Grid")
        props=layers["Precipitation"][0]["properties"]
        self.assertEqual(props["R_February_Mean"],58)
        self.assertEqual(props["Coverage_Fraction"],2/12)
        self.assertIsNone(props["R_Annual_Mean"])
        self.assertEqual(daily["Precipitation"][1],[])

    def test_individual_units_preserve_temperature_differences(self):
        point=make_point("p","Point",31,30);start,end=dt.date(2024,1,1),dt.date(2024,12,31)
        raw={p:{start+dt.timedelta(days=i):value for i in range(366)} for p,value in {"TS":28,"T2M":25}.items()}
        result=[dict(station=point,series=raw,status="OK",error="")]
        for unit,expected in (("Fahrenheit",82.4),("Kelvin",301.15)):
            layers,_=climate.prepare_outputs(result,["Land Surface Temperature"],start,end,"NASA POWER",unit_options={"unit_temp":unit})
            props=layers["Land Surface Temperature"][0]["properties"]
            self.assertAlmostEqual(props["LST_Annual_Mean"],expected)
            self.assertEqual(props["LST_minus_T_Annual"],5.4 if unit=="Fahrenheit" else 3)
            self.assertEqual(props["Measurement_Unit"],unit)

    def test_classification_outputs_highest_class_one_and_nodata_zero(self):
        import numpy as np
        from mirsad_desktop.analysis import classify_surface
        values=np.array([[1,2,3],[100,101,np.nan]])
        for method in ("Natural Breaks (Jenks)","Equal Interval","Quantile"):
            raster,info=classify_surface(values,method,2)
            self.assertEqual(raster[0,0],2);self.assertEqual(raster[1,1],1);self.assertEqual(raster[1,2],0)
            self.assertEqual(info["classes"],2)

    def test_regional_incomplete_dependencies_report_partial_and_keep_exports(self):
        from mirsad_desktop.regional import run_regional_project
        from shapely.geometry import box,mapping
        def response(session,url,query,cancelled):
            parameter=query["parameters"]
            return {"features":[{"geometry":{"coordinates":[31+dx,30+dy]},"properties":{"parameter":{
                parameter:{"202401":20} if parameter=="T2M" else {}}}}
                for dx,dy in ((0,0),(.1,0),(0,.1))]},url
        with tempfile.TemporaryDirectory(dir=ROOT/".test-output") as tmp,patch("mirsad_desktop.regional.request_json",response):
            output,manifest=run_regional_project(tmp,"Missing",{},mapping(box(30.5,29.5,31.5,30.5)),None,
                dt.date(2024,1,1),dt.date(2024,1,31),["Temperature"])
            self.assertEqual(manifest["status"],"partial")
            self.assertTrue(manifest["errors"])
            self.assertTrue((output/"Project.gdb").exists());self.assertTrue((output/"Temperature.xlsx").exists())

    def test_full_worker_climate_options_generate_monthly_and_classified_rasters(self):
        from mirsad_desktop.jobs import execute
        from openpyxl import load_workbook
        start,end=dt.date(2024,1,1),dt.date(2024,12,31)
        points=[make_point(str(i),"Point",31+dx,30+dy) for i,(dx,dy) in enumerate(((0,0),(.02,0),(0,.02),(.02,.02)))]
        results=[dict(station=p,series={key:{start+dt.timedelta(days=j):value for j in range(366)} for key,value in
                    {"T2M":20+i,"T2M_MAX":25+i,"T2M_MIN":15+i}.items()},status="OK",error="") for i,p in enumerate(points)]
        (ROOT/".test-output").mkdir(exist_ok=True)
        with tempfile.TemporaryDirectory(dir=ROOT/".test-output") as tmp,patch.object(climate,"download",return_value=(results,[],[])) as download:
            request=dict(operation="climate",points=points,folder=tmp,name="Options",metadata={},options=dict(start=str(start),end=str(end),
                         provider="NASA POWER",om_model="ERA5-Land (Highest Resolution)",modules=["Temperature"],monthly_rasters=True,
                         scopes=["Annual","Winter"],save_daily=False,raster_options=dict(crs="EPSG:32636",cell_size=1000,pdf=False,
                         fields=["T_Annual_Mean","T_January_Mean"],reclass_enabled=True,reclass_method="Equal Interval",reclass_classes=3)))
            result=execute(request,Path(tmp),lambda v:None,lambda:False)
            self.assertEqual(result["manifest"]["status"],"complete",result["manifest"]["errors"])
            self.assertEqual(download.call_args.kwargs["om_model"],"ERA5-Land (Highest Resolution)")
            output=Path(result["output"])
            self.assertEqual(len(list((output/"Atlas").glob("*.tif"))),4)
            book=load_workbook(output/"Temperature.xlsx",read_only=True)
            try:self.assertIn("Month",book.sheetnames);self.assertEqual(book["Month"].max_column,16)
            finally:book.close()


if __name__=="__main__":unittest.main()

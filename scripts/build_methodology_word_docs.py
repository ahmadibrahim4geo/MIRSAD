# -*- coding: utf-8 -*-
"""Generate Word (.docx) documents for Methodology Guide (06_Methodology_Adding_New_Climate_Module)."""
import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TARGET_DIR = os.path.join(BASE_DIR, "Documents", "06_Methodology_Adding_New_Climate_Module")
AR_MD = os.path.join(TARGET_DIR, "Adding_New_Climate_Module_Guide_AR.md")
EN_MD = os.path.join(TARGET_DIR, "Adding_New_Climate_Module_Guide_EN.md")
AR_DOCX = os.path.join(TARGET_DIR, "Adding_New_Climate_Module_Guide_AR.docx")
EN_DOCX = os.path.join(TARGET_DIR, "Adding_New_Climate_Module_Guide_EN.docx")

def convert_md_to_docx(src_md, dst_docx, is_arabic=False):
    doc = docx.Document()
    for s in doc.sections:
        s.top_margin = Inches(0.75)
        s.bottom_margin = Inches(0.75)
        s.left_margin = Inches(0.75)
        s.right_margin = Inches(0.75)

    with open(src_md, "r", encoding="utf-8") as f:
        lines = f.readlines()

    in_code = False
    in_table = False
    table_lines = []

    def flush_table(t_lines):
        if not t_lines:
            return
        rows_data = []
        for tl in t_lines:
            if not tl.strip() or "---" in tl:
                continue
            parts = [c.strip() for c in tl.strip().strip("|").split("|")]
            rows_data.append(parts)
        if not rows_data:
            return
        num_cols = max(len(r) for r in rows_data)
        tbl = doc.add_table(rows=len(rows_data), cols=num_cols)
        tbl.autofit = True
        for r_idx, r_data in enumerate(rows_data):
            row = tbl.rows[r_idx]
            is_head = (r_idx == 0)
            for c_idx, val in enumerate(r_data):
                if c_idx < num_cols:
                    cell = row.cells[c_idx]
                    cell.text = val
                    if is_head:
                        tcPr = cell._tc.get_or_add_tcPr()
                        tcPr.append(parse_xml('<w:shd {} w:fill="0284C7"/>'.format(nsdecls("w"))))
                        for p in cell.paragraphs:
                            for r in p.runs:
                                r.font.color.rgb = RGBColor(255, 255, 255)
                                r.font.bold = True
                    else:
                        if r_idx % 2 == 1:
                            tcPr = cell._tc.get_or_add_tcPr()
                            tcPr.append(parse_xml('<w:shd {} w:fill="F1F5F9"/>'.format(nsdecls("w"))))
        doc.add_paragraph()

    for line in lines:
        raw = line.rstrip("\r\n")
        sline = raw.strip()

        if sline.startswith("```"):
            if in_code:
                in_code = False
            else:
                if in_table:
                    flush_table(table_lines)
                    table_lines = []
                    in_table = False
                in_code = True
            continue

        if in_code:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(raw)
            run.font.name = "Consolas"
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor(30, 41, 59)
            continue

        if sline.startswith("|") and sline.endswith("|"):
            in_table = True
            table_lines.append(sline)
            continue
        elif in_table:
            flush_table(table_lines)
            table_lines = []
            in_table = False

        if not sline:
            continue

        if sline.startswith("# "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(sline[2:])
            run.font.size = Pt(18)
            run.font.bold = True
            run.font.color.rgb = RGBColor(15, 23, 42)
        elif sline.startswith("## "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(sline[3:])
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = RGBColor(2, 132, 199)
        elif sline.startswith("### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(sline[4:])
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = RGBColor(30, 41, 59)
        elif sline.startswith("#### "):
            p = doc.add_paragraph()
            run = p.add_run(sline[5:])
            run.font.size = Pt(10.5)
            run.font.bold = True
        elif sline.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(sline[2:].replace("[!IMPORTANT]", "تنبيه هام / IMPORTANT:").replace("[!WARNING]", "تحذير / WARNING:"))
            run.font.size = Pt(9.5)
            run.font.italic = True
            run.font.color.rgb = RGBColor(71, 85, 105)
        elif sline.startswith("- ") or sline.startswith("* "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(sline[2:])
            run.font.size = Pt(10)
        elif sline == "---":
            p = doc.add_paragraph()
            run = p.add_run("__________________________________________________________________")
            run.font.color.rgb = RGBColor(203, 213, 225)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(sline)
            run.font.size = Pt(10)

    if in_table:
        flush_table(table_lines)

    doc.save(dst_docx)
    print("Generated Word document: {}".format(dst_docx))

if __name__ == "__main__":
    convert_md_to_docx(AR_MD, AR_DOCX, is_arabic=True)
    convert_md_to_docx(EN_MD, EN_DOCX, is_arabic=False)

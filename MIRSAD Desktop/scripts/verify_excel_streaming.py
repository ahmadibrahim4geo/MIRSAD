"""Real default-capacity XLSX boundary; no network, CSV, or million-row RAM list."""
import argparse
import json
import sys
import time
from pathlib import Path

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/"src"))


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--rows",type=int,default=1000001)
    parser.add_argument("--output",type=Path,default=ROOT/".test-output/excel-streaming/Boundary.xlsx")
    args=parser.parse_args()
    if not 1<=args.rows<=5000000:parser.error("Use 1..5,000,000 rows for this explicit stress fixture")
    from mirsad_desktop.exports import workbook
    from openpyxl import load_workbook
    started=time.monotonic()
    def rows():
        for i in range(args.rows):
            if i%100000==0:print("Produced",i,"daily rows",flush=True)
            yield ["2024-02-29",i]
    result=workbook(args.output,[("Source_Data",["Source"],[["Generated sequential fixture"]]),("Daily",["Date","ID"],rows())])
    read=0;sheet_count=0
    for path in result["files"]:
        book=load_workbook(path,read_only=True,data_only=True)
        try:
            for name in book.sheetnames:
                if name=="Daily" or name.startswith("Daily_"):
                    sheet_count+=1
                    iterator=iter(book[name].values);next(iterator)
                    for date,identifier in iterator:
                        if identifier!=read:raise AssertionError("Missing, duplicated or reordered record at "+str(read))
                        read+=1
                    print("Read back",read,"rows",flush=True)
        finally:book.close()
    if read!=args.rows:raise AssertionError("Row count changed")
    report=dict(passed=True,rows=read,sheets=sheet_count,files=[str(p) for p in result["files"]],
                seconds=round(time.monotonic()-started,3),scope="Real default-capacity row boundary; not a full 1000-point/30-year climate pipeline benchmark")
    args.output.with_suffix(".acceptance.json").write_text(json.dumps(report,indent=2)+"\n",encoding="utf-8")
    print(json.dumps(report))


if __name__=="__main__":main()

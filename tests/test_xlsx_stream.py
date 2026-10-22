# -*- coding: utf-8 -*-
from __future__ import unicode_literals
import os
import json
import shutil
import sys
import tempfile
import unittest
import xml.etree.ElementTree as ET
import zipfile

ROOT=os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0,ROOT)
from mirsad_core.xlsx_stream import write_workbook,RowSpool,ExcelWriteError,NS


def read_rows(path):
    with zipfile.ZipFile(path) as archive:
        sheets=ET.fromstring(archive.read("xl/workbook.xml")).find("{%s}sheets"%NS)
        result=[]
        for i,sheet in enumerate(sheets,1):
            xml=ET.fromstring(archive.read("xl/worksheets/sheet%d.xml"%i))
            rows=[]
            for row in xml.find("{%s}sheetData"%NS):
                cells=[]
                for cell in row:
                    value=cell.find("{%s}v"%NS)
                    cells.append(''.join(cell.itertext()) if cell.get("t")=="inlineStr" else float(value.text) if value is not None else None)
                rows.append(cells)
            result.append((sheet.get("name"),rows,xml))
        return result


class SharedExcelTests(unittest.TestCase):
    def setUp(self):
        folder=os.path.join(ROOT,"tests",".output");
        if not os.path.isdir(folder):os.makedirs(folder)
        self.folder=tempfile.mkdtemp(prefix="shared-xlsx-",dir=folder)
        self.path=os.path.join(self.folder,"Climate.xlsx")

    def tearDown(self):shutil.rmtree(self.folder)

    def test_literal_arabic_ids_nulls_dimensions_and_independent_reader(self):
        write_workbook(self.path,[("Data",["ID","Name","Value"],[["0001","القاهرة",0],["=1+1","النقطة",None]],True)])
        tables=read_rows(self.path)
        self.assertEqual(tables[0][1][1],["0001","القاهرة",0.])
        self.assertEqual(tables[0][1][2],["=1+1","النقطة"])
        self.assertEqual(tables[0][2].find("{%s}dimension"%NS).get("ref"),"A1:C3")
        self.assertEqual(tables[0][2].find("{%s}sheetViews"%NS)[0].get("rightToLeft"),"1")
        if sys.version_info[0]>=3:
            from openpyxl import load_workbook
            book=load_workbook(self.path,read_only=True,data_only=False)
            try:
                self.assertEqual(book["Data"].max_row,3)
                self.assertEqual(book["Data"].cell(3,1).value,"=1+1")
                self.assertEqual(book["Data"].cell(3,1).data_type,"s")
                self.assertIsNone(book["Data"].cell(3,3).value)
            finally:book.close()

    def test_sheet_file_parts_index_and_repeated_single_pass_metadata(self):
        result=write_workbook(self.path,[("Data",["Name"],iter([["metadata"]]),False),
            ("Daily",["Date","ID"],(["2024-02-29",i] for i in range(11)),False)],sheet_capacity=3,part_daily_rows=4)
        self.assertEqual(len(result["files"]),3)
        values=[]
        for path in result["files"]:
            tables=read_rows(path)
            self.assertEqual(tables[0][1],[['Name'],['metadata']])
            for name,rows,_ in tables:
                if name.startswith("Daily"):values.extend(int(r[1]) for r in rows[1:])
        self.assertEqual(values,list(range(11)))
        self.assertEqual(sum(r["Records"] for r in result["index"]),11)
        self.assertTrue(any(name=="Export_Index" for name,_,_ in read_rows(self.path)))
        self.assertFalse(any(name.endswith(".csv") for name in os.listdir(self.folder)))

    def test_spool_and_successful_parts_survive_later_failure(self):
        rows=RowSpool(memory_bytes=32)
        try:
            for i in range(9):rows.append([i])
            first,second=iter(rows),iter(rows)
            self.assertEqual(next(first),next(second))
            self.assertEqual(rows[-1],[8])
            self.assertEqual(rows[2:4],[[2],[3]])
            def broken():
                for row in rows:yield row
                raise ValueError("fixture interrupted")
            with self.assertRaises(ExcelWriteError) as caught:
                write_workbook(self.path,[("Daily",["ID"],broken(),False)],sheet_capacity=3,part_daily_rows=4)
            self.assertEqual(len(caught.exception.files),1)
            actual=[int(r[0]) for _,data,_ in read_rows(caught.exception.files[0]) for r in data[1:]]
            self.assertEqual(actual,[4,5,6,7])
        finally:rows.close()

    def test_limits_and_existing_files_are_preserved(self):
        with self.assertRaises(ValueError):write_workbook(self.path,[("Data",["ID"],[],False)],sheet_capacity=1048576)
        with open(self.path,"wb") as handle:handle.write(b"original")
        with self.assertRaises(IOError):write_workbook(self.path,[("Data",["ID"],[],False)])
        with open(self.path,"rb") as handle:self.assertEqual(handle.read(),b"original")

    def test_native_failure_report_handles_ascii_and_arabic_on_both_runtimes(self):
        sys.path.insert(0,os.path.join(ROOT,"scripts"))
        from verify_native_xlsx import write_report
        for reason in ("Native timeout","توقف الاختبار"):
            write_report(self.folder,dict(passed=False,error=reason,timeout_seconds=120))
            with open(os.path.join(self.folder,"acceptance.json"),"rb") as handle:
                self.assertEqual(json.loads(handle.read().decode("utf-8"))["error"],reason)


if __name__=="__main__":unittest.main()

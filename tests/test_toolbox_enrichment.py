# -*- coding: utf-8 -*-
"""Daily optional Toolbox contracts; actual native processing is a separate gate."""
from __future__ import unicode_literals
import copy
import datetime as dt
import os
import shutil
import sys
import tempfile
import unittest
import test_runtime_stability as support
from mirsad_core import toolbox_enrichment as e
from mirsad_core import hybrid_pipeline as h


def fixture(start=dt.date(2024, 2, 1), end=dt.date(2024, 2, 29)):
    point = dict(oid=7, station_id="001234", wban="00001", station_name="القاهرة", lat=30., lon=31., calculation_elevation=0.)
    values = dict(T2M_MAX=25., T2M_MIN=15., T2MDEW=10., RH2M_MAX=84., RH2M_MIN=63.,
                  VPD_SOURCE=1.2, ALLSKY_SFC_SW_DWN=22.07, WS10M_CONFIRMED=3., SUNSHINE_DURATION=33300.)
    series = dict((p, dict((str(start+dt.timedelta(days=i)), value) for i in range((end-start).days+1))) for p, value in values.items())
    source = dict(series=series, source_summary="Open-Meteo era5", provenance={}, processing_by_variable={},
                  station=dict(station_key="001234|00001|30|31"))
    return point, source


class ToolboxEnrichment(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="enrichment_", dir=os.path.join(support.ROOT, "tests", "logs"))
        self.fake = support.ArcPy()
        support.core.arcpy = support.stations.arcpy = support.stations._core.arcpy = self.fake
        self.start, self.end = dt.date(2024, 2, 1), dt.date(2024, 2, 29)

    def tearDown(self):
        shutil.rmtree(self.tmp)

    def test_four_interfaces_append_optional_defaults_and_preserve_original_positions(self):
        for cls, count in ((support.core.PowerClimateAtlasGenerator, 83), (support.core.SpatialClimateAtlasGeneratorAR, 83),
                           (support.stations.WeatherStationsAtlasGenerator, 90), (support.stations.WeatherStationsAtlasGeneratorAR, 90)):
            params = cls().getParameterInfo()
            self.assertEqual(len(params), count)
            self.assertEqual([p.name for p in params[-14:-3]], [e.parameter_name(name) for name in e.ORDER])
            self.assertTrue(all(p.value is False for p in params[-14:-2]))
            if cls in (support.core.SpatialClimateAtlasGeneratorAR, support.stations.WeatherStationsAtlasGeneratorAR):
                self.assertIn("الرطوبة", params[-14].displayName)

    def test_parent_daily_offline_and_composite_validation(self):
        tool = support.core.PowerClimateAtlasGenerator()
        params = tool.getParameterInfo(); pd = dict((p.name, p) for p in params)
        pd["Climate_Modules"].value = "Relative Humidity"
        pd["Temporal_Resolution"].value = "Daily"
        pd["Climate_Data_Source"].value = "Open-Meteo ERA5 (Standalone - No NASA Supplement)"
        pd["Supplementary_RHmax"].value = True
        tool.updateParameters(params)
        self.assertTrue(pd["Supplementary_RHmax"].enabled)
        self.assertFalse(pd["Supplementary_Trange"].enabled)
        self.assertTrue(pd["Supplementary_Rasters"].enabled)
        pd["Climate_Data_Source"].value = "NASA POWER"
        tool.updateParameters(params)
        self.assertFalse(pd["Supplementary_RHmax"].enabled)
        self.assertFalse(pd["Supplementary_RHmax"].value)
        with self.assertRaises(ValueError):e.validate(["RHmax"], ["Relative Humidity"], "Daily", False, "NASA POWER")
        for temporal, offline, modules in (("Monthly", False, ["Relative Humidity"]), ("Daily", True, ["Relative Humidity"]), ("Daily", False, ["Wind"])):
            with self.assertRaises(ValueError):e.validate(["RHmax"], modules, temporal, offline, "NASA POWER")
        from mirsad_core.climate_products import ENABLED, descriptor
        monthly = next(name for name in ENABLED if descriptor(name)["temporal"] != "Daily")
        with self.assertRaises(ValueError):e.validate(["RHmax"], ["Relative Humidity"], "Daily", False, monthly)

    def test_all_variables_canonical_units_complete_month_zero_and_no_mutation(self):
        point, source = fixture(); original = copy.deepcopy((point, source))
        result = e.calculate(point, source, e.ORDER, self.start, self.end)
        f = result["fields"]
        self.assertEqual(f["RHmax_February_Mean"], 84.)
        self.assertEqual(f["Trange_February_Mean"], 10.)
        self.assertEqual(f["Sunshine_February_Mean"], 9.25)
        self.assertAlmostEqual(f["U2_February_Mean"], 2.2438532255, places=7)
        self.assertEqual(f["Rn_FAO56_Valid_Days"], 29)
        self.assertIsNone(f["Trange_Annual_Mean"])
        self.assertEqual(result["provenance"]["Rn_FAO56"][0]["elevation_m"], 0.)
        self.assertEqual((point, source), original)

    def test_unknown_elevation_and_invalid_hourly_dependencies_stay_null(self):
        point, source = fixture(); point.pop("calculation_elevation")
        source["processing_by_variable"] = {"T2MDEW": {"2024-02-29": dict(status="incomplete", valid_samples=23, expected_samples=24)}}
        source["series"]["RH2M_MAX"].pop("2024-02-29")
        original = copy.deepcopy(source)
        result = e.calculate(point, source, ["ea", "Rn_FAO56"], self.start, self.end)
        self.assertEqual(result["fields"]["ea_Valid_Days"], 28)
        self.assertIsNone(result["fields"]["ea_February_Mean"])
        self.assertEqual(result["fields"]["Rn_FAO56_Valid_Days"], 0)
        self.assertEqual(source, original)

    def test_input_identity_arabic_and_known_elevation_flags_survive_both_hosts(self):
        values = dict(OBJECTID=1, Station_ID="001234", WBAN="00001", Name="القاهرة", Country="EG", Elev_m=0.,
                      Point_Type="Ground Station", Z_Known=1)
        values["SHAPE@XY"] = (31., 30.)
        self.fake.sr.factoryCode = 4326
        self.fake.ListFields = lambda path: [support.Bag(name=name) for name in values]
        self.fake.da.SearchCursor = lambda path, fields: support.Cursor([tuple(values.get(field) for field in fields)])
        tool = support.core.PowerClimateAtlasGenerator();tool._supplementary_selected = ["Rn_FAO56"]
        point = tool._extract_points("fixture", self.fake.sr, support.Messages())[0]
        self.assertEqual((point["station_id"], point["wban"], point["station_name"]), ("001234", "00001", "القاهرة"))
        self.assertEqual(point["calculation_elevation"], 0.)
        values["Z_Known"] = 0
        point = tool._extract_points("fixture", self.fake.sr, support.Messages())[0]
        self.assertIsNone(point["calculation_elevation"])
        station_tool = support.stations.WeatherStationsAtlasGenerator();station_tool._supplementary_selected = ["Rn_FAO56"]
        values["Elev_m"] = None
        records = station_tool._load_all_stations("fixture")
        self.assertEqual(records[0]["WBAN"], "00001")
        self.assertFalse(records[0]["Z_Known"])
        station_tool._supplementary_selected=[];station_tool._daily_reference_et=False
        records=station_tool._load_all_stations("fixture")
        self.assertEqual(len(records),1)
        self.assertEqual(records[0]["Station_ID"],"001234")
        self.assertFalse(records[0]["Z_Known"])

    def test_summary_daily_weighting_and_interannual_ranges(self):
        start, end = dt.date(2023, 1, 1), dt.date(2024, 12, 31)
        values = dict((start+dt.timedelta(days=i), 10. if (start+dt.timedelta(days=i)).year == 2023 else 30.) for i in range(731))
        fields = e.summarize(values, start, end, "Test", ranges=True)
        self.assertEqual(fields["Test_Annual_Mean"], 20.)
        self.assertEqual(fields["Test_February_Range"], 20.)
        self.assertEqual(fields["Test_Annual_Range"], 0.)
        values.pop(dt.date(2024, 2, 29))
        missing = e.summarize(values, start, end, "Test", ranges=True)
        self.assertEqual(missing["Test_Annual_Mean"], 10.)
        self.assertIsNone(missing["Test_February_Range"])
        self.assertEqual(missing["Test_Valid_Days"], 730)

    def test_download_uses_selected_policy_deduplicates_locations_and_skips_network_for_ra(self):
        point, source = fixture(); calls = []
        def fetch(provider, station, start, end, required, model, session, cancelled):
            calls.append(provider)
            return dict(series=dict((p, source["series"].get(p, {})) for p in required), warnings=[], urls=[], state="COMPLETE")
        for provider, expected in (("NASA POWER", ["NASA POWER"]), ("Open-Meteo ERA5 (Standalone - No NASA Supplement)", ["Open-Meteo"])):
            calls[:] = []
            records = e.download([point, point], ["Trange"], self.start, self.end, provider, "era5", None, lambda event: None, lambda: False, fetch=fetch)
            self.assertEqual(calls, expected)
            self.assertEqual(len(records), 2)
            self.assertIs(records[0], records[1])
        calls[:] = []
        e.download([point], ["Ra"], self.start, self.end, "NASA POWER", "era5", None, lambda event: None, lambda: False, fetch=fetch)
        self.assertFalse(calls)

    def test_dictionary_shapefile_names_and_monthly_stretches_never_mix_units(self):
        grouped = e.fields_by_module(e.ORDER, self.start, self.end, ranges=True)
        all_fields = [f for fields in grouped.values() for f in fields]
        short = [e.short_field(f) for f in all_fields]
        self.assertEqual(len(short), len(set(short)))
        self.assertTrue(all(len(f) <= 10 for f in short))
        self.assertFalse(set(short) & set(support.core.SHP_FIELD_MAP.values()))
        defs = support.core.metadata_rows_for_modules(["Relative Humidity"], grouped)
        ea = next(d for d in defs if d["Field_Name"] == "ea_February_Mean")
        self.assertEqual(ea["Unit"], "kPa")
        self.assertNotEqual(support.core.monthly_group_key("Relative Humidity", "ea_February_Mean"),
                            support.core.monthly_group_key("Relative Humidity", "RHmax_February_Mean"))
        # The production QA must recognize the same optional SHP mappings
        # even when this run only requested a geodatabase.
        names = [field[0] for field in support.core.ADMIN_FIELDS] + grouped["Relative Humidity"]
        self.fake.ListFields = lambda path: [support.Bag(name=name) for name in names]
        checks = support.core.PowerClimateAtlasGenerator()._run_qa({"Relative Humidity": "fixture"},
            dict(vec=self.tmp), ["Relative Humidity"], [], [], lambda text: None, lambda text: None,
            wanted_fields_by_module={"Relative Humidity": grouped["Relative Humidity"]})
        self.assertTrue(next(ok for label, ok, detail in checks if label == "shapefile name length<=10 (or mapped)"))

    def test_audit_keeps_rejected_23_hour_day_and_year_filter(self):
        point, source = fixture()
        info = dict(status="incomplete", valid_samples=23, expected_samples=24, timezone="UTC", operation="max", input_temporal="hourly", output_temporal="daily", unit="percent")
        source["series"]["RH2M_MAX"].pop("2024-02-29")
        source["processing_attempts"] = [dict(source="Open-Meteo", variables={"RH2M_MAX": {"2024-02-29": info, "2023-02-28": info}})]
        point["supplementary"] = e.calculate(point, source, ["RHmax"], self.start, self.end)
        sheets = e.sheets([point], "Relative Humidity", ["RHmax"], self.start, self.end, True)
        try:
            audit = next(sheet for sheet in sheets if sheet[0] == "Hourly_Processing")
            rows = list(audit[2]);self.assertEqual(len(rows), 1)
            self.assertEqual(rows[0][9:11], [23, 24])
            self.assertFalse(rows[0][14])
            self.assertEqual(len(next(sheet for sheet in sheets if sheet[0] == "Supplementary_Daily")[2]), 28)
        finally:
            for sheet in sheets:
                if hasattr(sheet[2], "close"):sheet[2].close()

    @unittest.skipIf(sys.version_info[0] < 3, "Independent Desktop runtime uses Python 3")
    def test_desktop_and_toolbox_summaries_match_for_varied_daily_values_and_missingness(self):
        sys.path.insert(0, os.path.join(support.ROOT, "MIRSAD Desktop", "src"))
        from mirsad_desktop.enrichment import summarize as desktop_summary
        start, end = dt.date(2023, 1, 1), dt.date(2024, 12, 31)
        values = dict((start+dt.timedelta(days=i), float(i % 37)) for i in range(731))
        values.pop(dt.date(2024, 2, 29))
        for scopes in (None, ["Annual"], ["Winter", "Summer"], ["Seasonal"]):
            self.assertEqual(e.summarize(values, start, end, "Check", scopes, True),
                             desktop_summary(values, start, end, "Check", scopes, True))

    @unittest.skipIf(sys.version_info[0] < 3, "Independent Desktop runtime uses Python 3")
    def test_legacy_export_reads_real_xlsx_and_yearly_values_without_changing_raw(self):
        import openpyxl
        tool = support.core.PowerClimateAtlasGenerator()
        point, source = fixture(dt.date(2023, 2, 1), self.end)
        # Distinct complete Februaries prove each year is recomputed.
        for date in source["series"]["T2M_MAX"]:
            source["series"]["T2M_MAX"][date] = 25. if date.startswith("2023") else 35.
        point.update(monthly={"T2M": {(2023, 2): 20., (2024, 2): 30.}}, daily_raw={"T2M_MAX": {(2023, 2, 28): 25., (2024, 2, 29): 35.}}, fields={}, status="OK")
        start = dt.date(2023, 2, 1)
        point["supplementary"] = e.calculate(point, source, ["Trange"], start, self.end)
        before = copy.deepcopy(point)
        meta = dict(y0=2023, y1=2024, temporal="Daily", provider="NASA POWER", gap_method="none",
                    request_start=str(start), request_end=str(self.end), supplementary_start=str(start), supplementary_end=str(self.end),
                    supplementary_selected=["Trange"], supplementary_save_daily=True)
        fields = ["OBJECTID", "Station_ID", "Trange_February_Mean"]
        warnings = []
        tool._export_yearly_excel([point], [2023, 2024], "Temperature", self.tmp, meta, lambda text: None, warnings.append,
                                wanted_fields=["Trange_February_Mean"], save_daily=True)
        self.assertFalse(warnings, warnings)
        for year, expected in ((2023, 10.), (2024, 20.)):
            wb = openpyxl.load_workbook(os.path.join(self.tmp, "Temperature_%d.xlsx" % year), read_only=True)
            try:
                data = list(wb["Month"].values)
                self.assertEqual(data[1][data[0].index("Trange_February_Mean")], expected)
                self.assertIn("Daily", wb.sheetnames)
                self.assertTrue(all(str(row[3]).startswith(str(year)) for row in wb["Supplementary_Daily"].iter_rows(min_row=2, values_only=True)))
                self.assertIn("Field_Dictionary", wb.sheetnames)
            finally:wb.close()
        self.assertEqual(point, before)


if __name__ == "__main__":unittest.main()

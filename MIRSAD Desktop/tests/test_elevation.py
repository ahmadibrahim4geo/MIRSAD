"""DEM sampling, batched service behavior and real point export preservation."""
import copy
import json
import sys
import tempfile
import unittest
from pathlib import Path

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/"src"))
from mirsad_desktop.inputs import make_point
from mirsad_desktop.elevation import enrich_points,OPEN_METEO


class ElevationTests(unittest.TestCase):
    def setUp(self):
        import numpy as np
        import rasterio
        from rasterio.transform import from_origin
        (ROOT/".test-output").mkdir(exist_ok=True)
        self.tmp=tempfile.TemporaryDirectory(dir=ROOT/".test-output");self.addCleanup(self.tmp.cleanup)
        self.folder=Path(self.tmp.name);self.path=self.folder/"terrain.tif"
        data=np.array([[0,10,20,30],[40,50,-9999,70],[80,90,100,110],[120,130,140,150]],dtype="float32")
        with rasterio.open(self.path,"w",driver="GTiff",width=4,height=4,count=1,dtype="float32",crs="EPSG:4326",
                           transform=from_origin(30,31,1,1),nodata=-9999) as dst:dst.write(data,1)
        self.options=dict(source="local",path=str(self.path),surface_type="DTM",vertical_datum="Test datum; user declaration")

    def test_local_zero_nodata_outside_and_original_station_height(self):
        points=[make_point(str(i),"site",x,y,elevation=123) for i,(x,y) in enumerate(((30.5,30.5),(31.5,29.5),(32.5,29.5),(35,35)))]
        original=copy.deepcopy(points)
        enriched,report=enrich_points(points,self.options)
        self.assertEqual(points,original)
        self.assertEqual([p["dem_elevation"] for p in enriched],[0,50,None,None])
        self.assertEqual([p["elevation"] for p in enriched],[123]*4)
        self.assertEqual(enriched[0]["dem_source"],"Local DEM: terrain")
        self.assertIsNone(enriched[2]["dem_source"])
        self.assertEqual(report["failed"],2)
        self.assertEqual(report["points"][0]["surface_type"],"DTM")

    def test_bilinear_requires_all_four_neighbours_and_real_crs(self):
        points=[make_point("mid","site",31,30),make_point("hole","site",32,30)]
        enriched,_=enrich_points(points,dict(self.options,sampling="bilinear"))
        self.assertEqual(enriched[0]["dem_elevation"],25)
        self.assertIsNone(enriched[1]["dem_elevation"])
        with self.assertRaisesRegex(ValueError,"sampling"):
            enrich_points(points,dict(self.options,sampling="unknown"))

    def test_field_collision_and_cancellation_do_not_modify_input(self):
        point=make_point("id","site",30.5,30.5,attributes={"z_ELEV_m":42})
        with self.assertRaisesRegex(ValueError,"collision"):enrich_points([point],self.options)
        point=make_point("id","site",30.5,30.5)
        with self.assertRaises(InterruptedError):enrich_points([point],self.options,cancelled=lambda:True)
        self.assertNotIn("dem_elevation",point)

    def test_worker_export_writes_the_same_z_and_source_to_both_databases_and_excel(self):
        from mirsad_desktop.jobs import execute
        from openpyxl import load_workbook
        import pyogrio
        point=make_point("0001","النقطة",30.5,30.5,elevation=123)
        request=dict(operation="export",points=[point],folder=str(self.folder),name="Project",metadata={},
                     options=dict(crs="EPSG:32636",elevation_options=self.options))
        result=execute(request,self.folder,lambda e:None,lambda:False)
        self.assertEqual(result["manifest"]["status"],"complete",result["manifest"]["errors"])
        for filename in ("Project.gpkg","Project.gdb"):
            meta,_,_,data=pyogrio.raw.read(Path(result["output"])/filename,layer="Stations")
            row={k:v[0] for k,v in zip(meta["fields"],data)}
            self.assertEqual(row["Z_Elev_m"],0)
            self.assertEqual(row["Z_Source"],"Local DEM: terrain")
            self.assertEqual(row["Elevation_m"],123)
        book=load_workbook(Path(result["output"])/"Stations.xlsx",read_only=True,data_only=True)
        try:
            headers,row=list(book["Data"].values)[:2];row=dict(zip(headers,row))
            self.assertEqual(row["Z_Elev_m"],0)
            self.assertEqual(row["Elevation_m"],123)
            source_headers,source_row=list(book["Elevation_Source"].values)[:2]
            source_row=dict(zip(source_headers,source_row))
            self.assertEqual(source_row["Surface_Type"],"DTM")
            self.assertEqual(source_row["Z_Elev_m"],0)
            self.assertEqual(source_row["Z_Source"],"Local DEM: terrain")
            self.assertIn("Elevation_Fields",book.sheetnames)
        finally:book.close()
        self.assertNotIn("dem_elevation",point)

    def test_web_batch_limit_failure_isolation_and_null_source(self):
        points=[make_point(str(i),"site",31+i*.00001,30) for i in range(201)]
        class Response:
            def __init__(self,count,fail):self.count=count;self.fail=fail;self.headers={};self.url="https://example.invalid/elevation-fixture"
            def __enter__(self):return self
            def __exit__(self,*args):pass
            def raise_for_status(self):
                if self.fail:raise OSError("service unavailable")
            def iter_content(self,size):yield json.dumps(dict(elevation=[0]*self.count)).encode()
        class Session:
            def __init__(self):self.calls=[]
            def get(self,url,params,**kwargs):
                count=len(params["latitude"].split(","));self.calls.append(count)
                return Response(count,len(self.calls)==2)
        session=Session()
        enriched,report=enrich_points(points,dict(source="open-meteo"),session=session)
        self.assertEqual(session.calls,[100,100,1])
        self.assertEqual(report["failed"],100)
        self.assertEqual(enriched[0]["dem_source"],OPEN_METEO)
        self.assertIsNone(enriched[100]["dem_source"])
        self.assertEqual(enriched[-1]["dem_elevation"],0)

    def test_cache_keeps_real_zero_and_corrupt_entries_are_refetched(self):
        from unittest.mock import patch
        point=make_point("id","site",31,30)
        cache=self.folder/"cache"
        with patch("mirsad_desktop.climate.request_json",return_value=({"elevation":[0]},"fixture")) as request:
            first,_=enrich_points([point],dict(source="open-meteo"),cache_folder=cache)
            second,_=enrich_points([point],dict(source="open-meteo"),cache_folder=cache)
            self.assertEqual(request.call_count,1)
            self.assertEqual(second[0]["dem_elevation"],0)
            next(cache.glob("*.json")).write_text("broken",encoding="utf-8")
            enrich_points([point],dict(source="open-meteo"),cache_folder=cache)
            self.assertEqual(request.call_count,2)


if __name__=="__main__":unittest.main()

# -*- coding: utf-8 -*-
"""
Master Climate Atlas Classification & Specification Workbook Generator.
Builds the comprehensive professional 22-sheet cartographic classification guide:
  1. 01_عن_الدليل_About_Guide
  2. 02_حقول_الإدارة_Admin_Fields
  3. 03_المعايير_القياسية_Standards
  4..22: 19 dedicated sheets, one for each Climate Module (01_Temperature to 19_Trends_And_Anomalies)

Inside each of the 19 module sheets, indicators are organized into 4 tiered tables:
  • Level 1: تصنيف 3 فئات (3 Classes Overview - Executive)
  • Level 2: تصنيف 5 فئات (5 Classes Standard - Golden Standard)
  • Level 3: تصنيف 7 فئات (7 Classes Detailed - Academic)
  • Level 4: تصنيف 9 فئات (9 Classes High-Resolution - Wide Range)
Separated by distinct, stylish section divider banners.
No collapsed classes. Adaptive element-aware decimal precision. Global standard preservation.
"""
from __future__ import print_function
import os
import sys
import math
import shutil
import xml.etree.ElementTree as ET

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except Exception:
    HAS_OPENPYXL = False

SCRIPT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DEFAULT_BASE_DIR = r"C:\Users\ahmad\Desktop\Egypt Climate Data 1996-2025"


def find_excel_template():
    candidates = [
        os.path.join(SCRIPT_DIR, "Documents", "03_Excel_Dictionaries", "Climate_Atlas_Fields_Dictionary_AR_EN_Units.xlsx"),
        os.path.join(SCRIPT_DIR, "Documents", "Climate_Atlas_Fields_Dictionary_AR_EN_Units.xlsx"),
        os.path.join(SCRIPT_DIR, "docs", "03_Excel_Dictionaries", "Climate_Atlas_Fields_Dictionary_AR_EN_Units.xlsx"),
        os.path.join(SCRIPT_DIR, "docs", "Climate_Atlas_Fields_Dictionary_AR_EN_Units.xlsx"),
        os.path.join(SCRIPT_DIR, "Documents", "Fields_AR_EN_Units.xlsx"),
        os.path.join(SCRIPT_DIR, "Fields_AR_EN_Units.xlsx"),
        r"D:\My Software\NASA POWER Climate Atlas Generator\Documents\03_Excel_Dictionaries\Climate_Atlas_Fields_Dictionary_AR_EN_Units.xlsx",
        r"D:\My Software\NASA POWER Climate Atlas Generator\docs\03_Excel_Dictionaries\Climate_Atlas_Fields_Dictionary_AR_EN_Units.xlsx",
        r"D:\My Software\NASA POWER Climate Atlas Generator\Fields_AR_EN_Units.xlsx",
    ]
    for c in candidates:
        if os.path.exists(c):
            return c
    return candidates[0]


EXCEL_TEMPLATE = find_excel_template()

MODULE_FOLDER_MAP = {
    "Temperature": "01_Temperature",
    "01_Temperature": "01_Temperature",
    "Land Surface Temperature": "02_Land_Surface_Temperature",
    "Land_Surface_Temperature": "02_Land_Surface_Temperature",
    "02_Land_Surface_Temperature": "02_Land_Surface_Temperature",
    "19_Land_Surface_Temperature": "02_Land_Surface_Temperature",
    "Precipitation": "03_Precipitation",
    "02_Precipitation": "03_Precipitation",
    "03_Precipitation": "03_Precipitation",
    "Sea Level Pressure": "04_Sea_Level_Pressure",
    "Sea_Level_Pressure": "04_Sea_Level_Pressure",
    "03_Sea_Level_Pressure": "04_Sea_Level_Pressure",
    "04_Sea_Level_Pressure": "04_Sea_Level_Pressure",
    "Surface Pressure": "05_Surface_Pressure",
    "Surface_Pressure": "05_Surface_Pressure",
    "04_Surface_Pressure": "05_Surface_Pressure",
    "05_Surface_Pressure": "05_Surface_Pressure",
    "Wind": "06_Wind",
    "05_Wind": "06_Wind",
    "06_Wind": "06_Wind",
    "Relative Humidity": "07_Relative_Humidity",
    "Relative_Humidity": "07_Relative_Humidity",
    "06_Relative_Humidity": "07_Relative_Humidity",
    "07_Relative_Humidity": "07_Relative_Humidity",
    "Dew Point": "08_Dew_Point",
    "Dew_Point": "08_Dew_Point",
    "07_Dew_Point": "08_Dew_Point",
    "08_Dew_Point": "08_Dew_Point",
    "Solar Radiation": "09_Solar_Radiation",
    "Solar_Radiation": "09_Solar_Radiation",
    "08_Solar_Radiation": "09_Solar_Radiation",
    "09_Solar_Radiation": "09_Solar_Radiation",
    "UV Index": "10_UV_Index",
    "UV_Index": "10_UV_Index",
    "09_UV_Index": "10_UV_Index",
    "10_UV_Index": "10_UV_Index",
    "Cloud Cover": "11_Cloud_Cover",
    "Cloud_Cover": "11_Cloud_Cover",
    "10_Cloud_Cover": "11_Cloud_Cover",
    "11_Cloud_Cover": "11_Cloud_Cover",
    "Thermal Comfort": "12_Heat_Index",
    "Heat Index": "12_Heat_Index",
    "Heat_Index": "12_Heat_Index",
    "11_Heat_Index": "12_Heat_Index",
    "12_Heat_Index": "12_Heat_Index",
    "Wind Chill": "13_Wind_Chill",
    "Wind_Chill": "13_Wind_Chill",
    "12_Wind_Chill": "13_Wind_Chill",
    "13_Wind_Chill": "13_Wind_Chill",
    "Aridity & Drought": "14_De_Martonne_Aridity",
    "De Martonne Aridity": "14_De_Martonne_Aridity",
    "De_Martonne_Aridity": "14_De_Martonne_Aridity",
    "13_De_Martonne_Aridity": "14_De_Martonne_Aridity",
    "14_De_Martonne_Aridity": "14_De_Martonne_Aridity",
    "Evapotranspiration": "15_Evapotranspiration",
    "14_Evapotranspiration": "15_Evapotranspiration",
    "15_Evapotranspiration": "15_Evapotranspiration",
    "UNEP Aridity": "16_UNEP_Aridity",
    "UNEP_Aridity": "16_UNEP_Aridity",
    "15_UNEP_Aridity": "16_UNEP_Aridity",
    "16_UNEP_Aridity": "16_UNEP_Aridity",
    "Water Deficit": "17_Water_Deficit",
    "Water_Deficit": "17_Water_Deficit",
    "16_Water_Deficit": "17_Water_Deficit",
    "17_Water_Deficit": "17_Water_Deficit",
    "Dry Months": "18_Dry_Months",
    "Dry_Months": "18_Dry_Months",
    "17_Dry_Months": "18_Dry_Months",
    "18_Dry_Months": "18_Dry_Months",
    "Trends & Anomalies": "19_Trends_And_Anomalies",
    "Trends_And_Anomalies": "19_Trends_And_Anomalies",
    "18_Trends_And_Anomalies": "19_Trends_And_Anomalies",
    "19_Trends_And_Anomalies": "19_Trends_And_Anomalies",
}

MODULES_METADATA = [
    {
        "code": "01_Temperature",
        "pfx": "01",
        "name_ar": u"درجة الحرارة",
        "name_en": "Temperature",
        "unit": "°C",
        "folder": "01_Temperature",
        "fc": "01_Temperature",
        "desc_ar": u"المتوسطات الحسابية لدرجات الحرارة T2M وفق معايير WMO للأطلس المناخي (سنوي، فصلي، ونهايات عظمى وصغرى).",
        "method_en": "NASA POWER T2M, T2M_MAX, T2M_MIN / WMO Climatological Normals Standard"
    },
    {
        "code": "02_Land_Surface_Temperature",
        "pfx": "02",
        "name_ar": u"حرارة سطح الأرض",
        "name_en": "Land Surface Temperature",
        "unit": "°C",
        "folder": "02_Land_Surface_Temperature",
        "fc": "02_Land_Surface_Temperature",
        "desc_ar": u"درجة حرارة سطح الأرض المقاسة بالاستشعار الفضائي (MERRA-2 TS) ومتوسطاتها وفروقها مع حرارة الهواء.",
        "method_en": "NASA POWER MERRA-2 Earth Skin Temperature (TS)"
    },
    {
        "code": "03_Precipitation",
        "pfx": "03",
        "name_ar": u"الأمطار والتساقط",
        "name_en": "Precipitation",
        "unit": "mm",
        "folder": "03_Precipitation",
        "fc": "03_Precipitation",
        "desc_ar": u"التساقط التراكمي PRECTOTCORR سنوياً وفصلياً وشهرياً، مع حساب المدى والتباين الفصلي.",
        "method_en": "NASA POWER PRECTOTCORR / Annual, Seasonal Totals & Cumulative Ranges"
    },
    {
        "code": "04_Sea_Level_Pressure",
        "pfx": "04",
        "name_ar": u"ضغط مستوى سطح البحر",
        "name_en": "Sea Level Pressure",
        "unit": "hPa",
        "folder": "04_Sea_Level_Pressure",
        "fc": "04_Sea_Level_Pressure",
        "desc_ar": u"متوسطات ضغط مستوى سطح البحر SLP المعدل طبوغرافياً وفق المعيار البارومتري الدولي لمقارنة الأنظمة الجوية.",
        "method_en": "NASA POWER SLP / Reduced to Standard Mean Sea Level (MSLP)"
    },
    {
        "code": "05_Surface_Pressure",
        "pfx": "05",
        "name_ar": u"الضغط السطحي الفعلي",
        "name_en": "Surface Pressure",
        "unit": "hPa",
        "folder": "05_Surface_Pressure",
        "fc": "05_Surface_Pressure",
        "desc_ar": u"متوسطات الضغط الجوي السطحي الحقيقي PS عند منسوب الارتفاع الطبوغرافي الفعلي للموقع الجغرافي.",
        "method_en": "NASA POWER PS / Actual Local Topographic Station Surface Pressure"
    },
    {
        "code": "06_Wind",
        "pfx": "06",
        "name_ar": u"الرياح السطحية (سرعة واتجاه)",
        "name_en": "Surface Wind (Speed & Direction)",
        "unit": "m/s, °",
        "folder": "06_Wind",
        "fc": "06_Wind",
        "desc_ar": u"متوسطات سرعة الرياح WS10M والمتوسط الشعاعي الدائري لاتجاه الرياح WD10M باستخدام دالة atan2(U,V).",
        "method_en": "NASA POWER WS10M, WD10M / Circular Mean Vector Trigonometry atan2(U,V)"
    },
    {
        "code": "07_Relative_Humidity",
        "pfx": "07",
        "name_ar": u"الرطوبة النسبية",
        "name_en": "Relative Humidity",
        "unit": "%",
        "folder": "07_Relative_Humidity",
        "fc": "07_Relative_Humidity",
        "desc_ar": u"نسبة بخار الماء الفعلي إلى التشبع RH2M عند درجة حرارة الهواء السطحي (سنوي، فصلي، شهري).",
        "method_en": "NASA POWER RH2M / WMO Psychrometric Standards"
    },
    {
        "code": "08_Dew_Point",
        "pfx": "08",
        "name_ar": u"نقطة الندى",
        "name_en": "Dew Point",
        "unit": "°C",
        "folder": "08_Dew_Point",
        "fc": "08_Dew_Point",
        "desc_ar": u"درجة الحرارة T2MDEW التي يتكثف عندها بخار الماء في الهواء الجوي تحت ضغط ثابت مع حساب عجز نقطة الندى.",
        "method_en": "NASA POWER T2MDEW / Magnus-Tetens Equation Standard"
    },
    {
        "code": "09_Solar_Radiation",
        "pfx": "09",
        "name_ar": u"الإشعاع الشمسي",
        "name_en": "Solar Radiation",
        "unit": "kWh/m²/day",
        "folder": "09_Solar_Radiation",
        "fc": "09_Solar_Radiation",
        "desc_ar": u"الإشعاع الشمسي الكلي الساقط على سطح أفقي ALLSKY_SFC_SW_DWN، مع المجموع السنوي التراكمي الساقط.",
        "method_en": "NASA POWER ALLSKY_SFC_SW_DWN / Total Downward Shortwave Flux"
    },
    {
        "code": "10_UV_Index",
        "pfx": "10",
        "name_ar": u"مؤشر الأشعة فوق البنفسجية",
        "name_en": "UV Index",
        "unit": "Index",
        "folder": "10_UV_Index",
        "fc": "10_UV_Index",
        "desc_ar": u"مؤشر الأشعة فوق البنفسجية القصوى السطحية لحماية الصحة العامة والتخطيط السياحي والبيئي وفق منظمة الصحة العالمية WHO.",
        "method_en": "WHO Global Solar UV Index Specification"
    },
    {
        "code": "11_Cloud_Cover",
        "pfx": "11",
        "name_ar": u"الغطاء السحابي",
        "name_en": "Cloud Cover",
        "unit": "%",
        "folder": "11_Cloud_Cover",
        "fc": "11_Cloud_Cover",
        "desc_ar": u"نسبة تغطية السحب الكلية في قبة السماء السنوية والفصلية والشهرية من استشعار CERES.",
        "method_en": "NASA CERES Total Cloud Area Fraction Standard"
    },
    {
        "code": "12_Heat_Index",
        "pfx": "12",
        "name_ar": u"الراحة الحرارية والرطوبة",
        "name_en": "Heat Index & Thermal Comfort",
        "unit": "°C",
        "folder": "12_Heat_Index",
        "fc": "12_Heat_Index",
        "desc_ar": u"مؤشرات الإجهاد والراحة الحرارية الناتجة عن تضافر الحرارة والرطوبة (Rothfusz Heat Index و Humidex).",
        "method_en": "Rothfusz (1990) NWS Heat Index & Masterton (1979) Humidex Standards"
    },
    {
        "code": "13_Wind_Chill",
        "pfx": "13",
        "name_ar": u"لسعة الرياح والبرودة",
        "name_en": "Wind Chill",
        "unit": "°C",
        "folder": "13_Wind_Chill",
        "fc": "13_Wind_Chill",
        "desc_ar": u"درجة الحرارة المحسوسة الفعالة الناتجة عن سرعة الرياح في الأجواء الباردة والشتوية.",
        "method_en": "Osczevski & Bluestein (2005) Wind Chill Standard"
    },
    {
        "code": "14_De_Martonne_Aridity",
        "pfx": "14",
        "name_ar": u"دليل دومارتون للجفاف",
        "name_en": "De Martonne Aridity Index",
        "unit": "Index",
        "folder": "14_De_Martonne_Aridity",
        "fc": "14_De_Martonne_Aridity",
        "desc_ar": u"دليل دومارتون الكلاسيكي للجفاف والمناخ الحيوي IDM = P / (T + 10).",
        "method_en": "De Martonne (1926) Aridity Index Standard"
    },
    {
        "code": "15_Evapotranspiration",
        "pfx": "15",
        "name_ar": u"البخر-نتح المرجعي",
        "name_en": "Reference Evapotranspiration",
        "unit": "mm",
        "folder": "15_Evapotranspiration",
        "fc": "15_Evapotranspiration",
        "desc_ar": u"معدل البخر-نتح المرجعي المحسوب وفق معادلة Hargreaves-Samani السنوية والفصلية والشهرية.",
        "method_en": "Hargreaves & Samani (1985) Reference Evapotranspiration FAO-56"
    },
    {
        "code": "16_UNEP_Aridity",
        "pfx": "16",
        "name_ar": u"مؤشر الجفاف العالمي UNEP",
        "name_en": "UNEP Aridity Index",
        "unit": "Index",
        "folder": "16_UNEP_Aridity",
        "fc": "16_UNEP_Aridity",
        "desc_ar": u"مؤشر الجفاف الدولي UNEP AI = P / PET لتصنيف الأقاليم الجافة وفق أطلس التصحر العالمي.",
        "method_en": "UNEP World Atlas of Desertification Standard"
    },
    {
        "code": "17_Water_Deficit",
        "pfx": "17",
        "name_ar": u"العجز المائي المناخي",
        "name_en": "Climatic Water Deficit",
        "unit": "mm",
        "folder": "17_Water_Deficit",
        "fc": "17_Water_Deficit",
        "desc_ar": u"العجز المائي الصافي (P - PET) لتقييم الموازنة المائية والاحتياجات المائية الزراعية.",
        "method_en": "Thornthwaite & Mather (1955) Climatic Water Balance"
    },
    {
        "code": "18_Dry_Months",
        "pfx": "18",
        "name_ar": u"عدد شهور الجفاف",
        "name_en": "Dry Months Count",
        "unit": "شهر (Months)",
        "folder": "18_Dry_Months",
        "fc": "18_Dry_Months",
        "desc_ar": u"عدد شهور السنة التي يقل فيها المطر عن ضعف درجة الحرارة (P < 2T) وفق معيار Bagnouls & Gaussen.",
        "method_en": "Bagnouls & Gaussen (1957) Xerothermic Index"
    },
    {
        "code": "19_Trends_And_Anomalies",
        "pfx": "19",
        "name_ar": u"اتجاهات وشذوذ المناخ",
        "name_en": "Climate Trends & Anomalies",
        "unit": "°C/dec, mm/dec, %, °C, mm",
        "folder": "19_Trends_And_Anomalies",
        "fc": "19_Trends_And_Anomalies",
        "desc_ar": u"اتجاهات التغير لعقد من الزمان (Decadal Trends) وشذوذ المناخ مقارنة بالمعدل الطبيعي 1991-2020.",
        "method_en": "WMO Climate Anomaly Standard & Least-Squares Linear Trends"
    }
]


def mod_field_matches(mod_str, pfx, code):
    norm = MODULE_FOLDER_MAP.get(str(mod_str).strip(), str(mod_str).strip())
    if norm == code or norm.startswith(pfx + "_"):
        return True
    return str(mod_str).strip().startswith(pfx + "_")


def get_clean_unit_symbol(unit, fld):
    u = str(unit or "").strip()
    if "Pct" in fld:
        return "%"
    if "Trend" in fld:
        return "°C/dec" if fld.startswith("T_") else "mm/dec"
    if "Anom" in fld:
        return "%" if "Pct" in fld else ("°C" if fld.startswith("T_") else "mm")
    if "Dry_Months" in fld:
        return "شهر"
    if "Aridity" in fld or "UV_" in fld:
        return "Index"
    if fld.startswith(("T_", "HI_", "WBGT_", "WC_", "TS_", "Td_", "LST_")):
        return "°C"
    if fld.startswith(("R_", "ET_", "PET_")) or "Deficit" in fld or "Evap" in fld:
        return "mm"
    if fld.startswith(("PS_", "PSL_")) or "Pressure" in fld:
        return "hPa"
    if fld.startswith("W_Spd"):
        return "m/s"
    if fld.startswith("W_Dir"):
        return "°"
    if fld.startswith(("RH_", "Cld_")):
        return "%"
    if fld.startswith("Sol_"):
        return "kWh/m²/year" if "Annual_Total" in fld else "kWh/m²/day"
    if u in ("°C", "°C/decade", "mm", "hPa", "m/s", "%", "Index"):
        return u
    return u if u and u != "-" else ""


def compute_smart_classes(full_f, vmin, vmax, n_classes, unit=""):
    """
    Computes analytical, strictly non-overlapping, non-collapsed class intervals
    directly from raster empirical min and max.
    Adheres strictly to Global Standards where applicable (WHO UV, UNEP, Directional Wind, NOAA HI),
    and applies adaptive dynamic decimal precision to prevent artificial integer truncation.
    Returns: (step_formatted_str, method_str, list_of_n_labels_from_highest_to_lowest, precision)
    """
    u_sym = get_clean_unit_symbol(unit, full_f)
    u_str = (" " + u_sym) if u_sym and not u_sym.startswith("°") and u_sym != "%" else (u_sym or "")

    # 1. Global Standard: UV Index (WHO Global Solar UV Index Standard)
    if full_f.startswith("UV_"):
        if n_classes == 5:
            labels = [
                u"11.0+ Index (شديد الخطورة - Extreme)",
                u"8.0 - 10.9 Index (مرتفع جداً - Very High)",
                u"6.0 - 7.9 Index (مرتفع - High)",
                u"3.0 - 5.9 Index (معتدل - Moderate)",
                u"0.0 - 2.9 Index (منخفض - Low)"
            ]
            return "فئات WHO القياسية", u"معيار منظمة الصحة العالمية (WHO UV Standard)", labels, 1
        elif n_classes == 3:
            labels = [
                u"8.0+ Index (مرتفع جداً إلى متطرف - Very High to Extreme)",
                u"3.0 - 7.9 Index (معتدل إلى مرتفع - Moderate to High)",
                u"0.0 - 2.9 Index (منخفض - Low)"
            ]
            return "فئات WHO المجمعة", u"معيار منظمة الصحة العالمية (WHO Consolidated)", labels, 1

    # 2. Global Standard: UNEP Aridity Index
    if full_f == "UNEP_Aridity_Annual" and n_classes == 5:
        labels = [
            u"> 0.65 (رطب - Humid)",
            u"0.50 - 0.65 (شبه رطب جاف - Dry sub-humid)",
            u"0.20 - 0.50 (شبه قاحل - Semi-arid)",
            u"0.05 - 0.20 (قاحل - Arid)",
            u"< 0.05 (شديد القحولة - Hyper-arid)"
        ]
        return "فئات UNEP القياسية", u"معيار أطلس التصحر الدولي (UNEP Standard)", labels, 3

    # 3. Wind Direction (Directional Azimuth Sectors)
    if full_f.startswith("W_Dir"):
        if n_classes == 5:
            labels = [
                u"288 - 360° (شمالي غربي إلى شمالي)",
                u"216 - 288° (غربي إلى شمالي غربي)",
                u"144 - 216° (جنوبي إلى جنوبي غربي)",
                u"72 - 144° (شرقي إلى جنوبي شرقي)",
                u"0 - 72° (شمالي إلى شمالي شرقي)"
            ]
            return "72°", u"قطاعات البوصلة الاتجاهية (Directional Azimuth)", labels, 0
        elif n_classes == 3:
            labels = [
                u"240 - 360° (غربي إلى شمالي)",
                u"120 - 240° (جنوبي إلى جنوبي غربي)",
                u"0 - 120° (شمالي إلى جنوبي شرقي)"
            ]
            return "120°", u"قطاعات البوصلة الاتجاهية (Directional Azimuth)", labels, 0

    # 4. Global Standard: Heat Index / Humidex (NOAA Danger categories)
    if full_f == "HI_Annual_Mean" and n_classes == 5 and vmax >= 40:
        labels = [
            u"≥ 54.0°C (خطر داهم - Extreme Danger)",
            u"41.0 - 53.9°C (خطر - Danger)",
            u"32.0 - 40.9°C (حذر شديد - Extreme Caution)",
            u"27.0 - 31.9°C (تنبيه - Caution)",
            u"< 27.0°C (طبيعي / مريح - Safe)"
        ]
        return "فئات NOAA للراحة الحرارية", u"معيار NOAA / NWS للراحة الحرارية", labels, 1

    # 5. General Analytical Equal Interval with Dynamic Precision
    raw_min = float(vmin)
    if abs(raw_min) < 1e-5:
        raw_min = 0.0
    raw_max = float(vmax)
    if abs(raw_max) < 1e-5:
        raw_max = 0.0
    span = raw_max - raw_min
    if span <= 0:
        span = 1.0

    raw_step = span / float(n_classes)

    # Determine precision based on step size and variable nature
    if raw_step < 0.05:
        prec = 3
    elif raw_step < 0.2:
        prec = 2
    elif raw_step < 2.0:
        prec = 1
    else:
        # Step >= 2.0
        if full_f.startswith(("R_", "ET_", "PET_", "PS_", "Sol_Annual_Total", "Water_Deficit")) and "Pct" not in full_f:
            prec = 0 if abs(raw_step - round(raw_step)) < 0.15 else 1
        elif full_f.startswith(("RH_", "Cld_")):
            prec = 0
        elif full_f.startswith(("T_", "HI_", "Td_", "WC_")):
            prec = 0 if (abs(raw_step - round(raw_step)) < 0.1 and abs(raw_min - round(raw_min)) < 0.1) else 1
        else:
            prec = 1

    # Format step string
    if prec == 0:
        step_str = "{}{}".format(int(round(raw_step)), u_str)
    else:
        step_str = "{:.{}f}{}".format(raw_step, prec, u_str)

    # Generate breaks
    breaks = [round(raw_min + i * raw_step, prec) for i in range(n_classes + 1)]
    breaks[0] = round(raw_min, prec)
    breaks[-1] = round(raw_max, prec)

    # Ensure strictly increasing breaks (prevent any collapsed classes)
    inc = round(10 ** (-prec), prec)
    for i in range(n_classes):
        if breaks[i + 1] <= breaks[i]:
            breaks[i + 1] = round(breaks[i] + inc, prec)

    # Build labels
    labels = []
    for i in range(n_classes):
        low = breaks[i]
        high = breaks[i + 1]
        if prec == 0:
            l_str = str(int(round(low)))
            h_str = str(int(round(high)))
        else:
            l_str = "{:.{}f}".format(low, prec)
            h_str = "{:.{}f}".format(high, prec)

        sep = u" إلى " if (l_str.startswith("-") or h_str.startswith("-")) else u" - "
        lbl = u"{}{}{}{}".format(l_str, sep, h_str, u_str)
        labels.append(lbl)

    ordered_labels = list(reversed(labels))
    method_str = u"فئات متساوية استرشادية (Equal Interval)"
    return step_str, method_str, ordered_labels, prec


# Legacy compatibility functions
def fmt_val(v, step=None):
    if v is None:
        return ""
    if abs(v - round(v)) < 1e-4:
        return str(int(round(v)))
    if abs(v * 10 - round(v * 10)) < 1e-4:
        return "{:.1f}".format(round(v, 1))
    if abs(v * 100 - round(v * 100)) < 1e-4:
        return "{:.2f}".format(round(v, 2))
    return "{:.2f}".format(v)


def get_arcmap_classes(fld, vmin, vmax, n_classes, precision=0, unit=""):
    step_str, method_str, labels, prec = compute_smart_classes(fld, vmin, vmax, n_classes, unit)
    span = float(vmax) - float(vmin)
    step_val = span / float(n_classes) if span > 0 else 1.0
    breaks = [float(vmin) + i * step_val for i in range(n_classes + 1)]
    return step_val, breaks, labels


def get_equal_classes(fld, vmin, vmax, n_classes, unit=""):
    return get_arcmap_classes(fld, vmin, vmax, n_classes, precision=0, unit=unit)[:2]


def get_equal_classes_decimal(fld, vmin, vmax, n_classes, unit=""):
    return get_arcmap_classes(fld, vmin, vmax, n_classes, precision=1, unit=unit)[:2]


def load_authoritative_fields():
    """Loads authoritative field metadata from Fields_AR_EN_Units.xlsx and/or PYT."""
    fields_list = []
    if os.path.exists(EXCEL_TEMPLATE):
        try:
            wb = openpyxl.load_workbook(EXCEL_TEMPLATE)
            if 'Climate_Fields' in wb.sheetnames:
                ws = wb['Climate_Fields']
                for r in range(2, ws.max_row + 1):
                    f_short = ws.cell(r, 2).value
                    f_full = ws.cell(r, 3).value
                    cmod = ws.cell(r, 4).value
                    desc_ar = ws.cell(r, 5).value
                    unit = ws.cell(r, 6).value
                    title_ar = ws.cell(r, 7).value
                    if f_full and 'Admin' not in str(cmod):
                        norm_mod = MODULE_FOLDER_MAP.get(str(cmod).strip(), str(cmod).strip())
                        fields_list.append({
                            'full_f': str(f_full).strip(),
                            'short_f': str(f_short or f_full).strip(),
                            'module': norm_mod,
                            'desc_ar': str(desc_ar or "").strip(),
                            'unit': str(unit or "").strip(),
                            'title_ar': str(title_ar or f_full).strip()
                        })
        except Exception:
            pass

    en_names = {}
    pyt_path = os.path.join(SCRIPT_DIR, "Toolboxes/MIRSAD_Climate_Atlas_Generator_10_8.pyt")
    if not os.path.exists(pyt_path):
        pyt_path = os.path.join(SCRIPT_DIR, "Toolboxes/POWER_Climate_Atlas_Generator_10_8.pyt")
    if os.path.exists(pyt_path):
        try:
            m = type(sys)("pyt_tmp")
            exec(compile(open(pyt_path, "rb").read(), "pyt_tmp", "exec"), m.__dict__)
            en_names = {f[0]: f[1] for f in getattr(m, "FIELD_DEFS", [])}
        except Exception:
            pass

    for f in fields_list:
        f['en_name'] = en_names.get(f['full_f'], f['full_f'].replace('_', ' '))

    return fields_list


def build_master_classification_workbook(base_dir=None, study_area_name="Egypt", target_excel=None, spatial_info=None, tech_info=None, raster_stats=None):
    r"""
    Builds the comprehensive 22-sheet Master Climate Atlas Classification & Specification Workbook:
      - Sheet 1: 01_عن_الدليل_About_Guide
      - Sheet 2: 02_حقول_الإدارة_Admin_Fields
      - Sheet 3: 03_المعايير_القياسية_Standards
      - Sheets 4..22: 19 dedicated sheets, one for each Climate Module (01_Temperature to 19_Trends_And_Anomalies).
    Each module sheet contains 4 tiered tables (3, 5, 7, 9 classes) separated by banners.
    Automatically saves copies to:
      1. target_excel (typically 00_Tables_And_Reports\Climate_Atlas_Classification_Guide.xlsx)
      2. base_dir\Climate_Atlas_Classification_Guide.xlsx (Root Output Workspace)
    """
    base_dir = base_dir or DEFAULT_BASE_DIR
    if not target_excel:
        target_excel = os.path.join(base_dir, "00_Tables_And_Reports", "Climate_Atlas_Classification_Guide.xlsx")

    # If running under Python 2.7 / ArcMap without openpyxl, delegate via subprocess to Python 3
    if not HAS_OPENPYXL:
        import subprocess
        import json
        import tempfile
        if os.environ.get("ATLAS_GUIDE_SUBPROCESS"):
            return target_excel
        try:
            payload = {
                "base_dir": base_dir,
                "study_area_name": study_area_name,
                "target_excel": target_excel,
                "spatial_info": spatial_info or {},
                "tech_info": tech_info or {},
                "raster_stats": raster_stats or {}
            }
            json_file = os.path.join(tempfile.gettempdir(), "atlas_guide_args.json")
            with open(json_file, "w") as jf:
                json.dump(payload, jf)
            script_path = os.path.abspath(__file__)
            env = dict(os.environ)
            env["ATLAS_GUIDE_SUBPROCESS"] = "1"
            py3_cmds = [
                ["python", script_path, "--json_args", json_file],
                ["py", "-3", script_path, "--json_args", json_file],
                [os.path.expandvars(r"%LOCALAPPDATA%\Microsoft\WindowsApps\python.exe"), script_path, "--json_args", json_file]
            ]
            for cmd in py3_cmds:
                try:
                    ret = subprocess.call(cmd, env=env)
                    if ret == 0 and os.path.exists(target_excel):
                        return target_excel
                except Exception:
                    pass
        except Exception:
            pass
        return target_excel

    # 1. Parse all raster aux.xml files or query ArcPy for exact empirical Min, Max, Mean and folder location
    if raster_stats is None:
        raster_stats = {}
    else:
        raster_stats = dict(raster_stats)

    if os.path.exists(base_dir):
        for root, dirs, files in os.walk(base_dir):
            dirs[:] = [d for d in dirs if not d.startswith('_') and not d.startswith('.') and 'retired' not in d.lower()]
            for f in sorted(files):
                if f.lower().endswith('.tif') and not f.lower().endswith('_cls.tif'):
                    rel_folder = os.path.relpath(root, base_dir)
                    tif_name = f
                    fld_name = f[:-4]
                    if fld_name in raster_stats and raster_stats[fld_name].get('min') is not None:
                        continue
                    xml_path = os.path.join(root, f + '.aux.xml')
                    s_min, s_max, s_mean = None, None, None
                    if os.path.exists(xml_path):
                        try:
                            tree = ET.parse(xml_path)
                            root_elem = tree.getroot()
                            for mdi in root_elem.findall('.//MDI'):
                                k = mdi.attrib.get('key')
                                if k == 'STATISTICS_MINIMUM': s_min = float(mdi.text)
                                elif k == 'STATISTICS_MAXIMUM': s_max = float(mdi.text)
                                elif k == 'STATISTICS_MEAN': s_mean = float(mdi.text)
                            if s_min is None:
                                hmin = root_elem.find('.//HistMin')
                                if hmin is not None and hmin.text: s_min = float(hmin.text)
                            if s_max is None:
                                hmax = root_elem.find('.//HistMax')
                                if hmax is not None and hmax.text: s_max = float(hmax.text)
                        except Exception:
                            pass
                    # ArcPy direct query fallback if statistics not found in aux.xml
                    if s_min is None or s_max is None:
                        try:
                            import arcpy
                            tif_fpath = os.path.join(root, f)
                            s_min = float(arcpy.GetRasterProperties_management(tif_fpath, "MINIMUM")[0])
                            s_max = float(arcpy.GetRasterProperties_management(tif_fpath, "MAXIMUM")[0])
                            s_mean = float(arcpy.GetRasterProperties_management(tif_fpath, "MEAN")[0])
                        except Exception:
                            pass
                    if s_min is not None and s_max is not None:
                        raster_stats[fld_name] = {
                            'folder': rel_folder,
                            'tif': tif_name,
                            'min': s_min,
                            'max': s_max,
                            'mean': s_mean if s_mean is not None else (s_min + s_max) / 2.0
                        }

    # Load authoritative climate fields
    climate_fields = load_authoritative_fields()

    # Styling Constants
    NAVY_HEADER = "1F4E79"
    HEADER_FONT = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    HEADER_FILL = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
    HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)

    THIN_BORDER = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    DATA_FONT = Font(name="Segoe UI", size=9)
    DATA_FONT_BOLD = Font(name="Segoe UI", size=9, bold=True)
    CODE_FONT = Font(name="Consolas", size=9)
    CODE_FONT_BOLD = Font(name="Consolas", size=9, bold=True)

    ROW_EVEN_FILL = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
    ROW_WHITE_FILL = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    SECTION_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

    sp_info = spatial_info or {}
    if (not sp_info or not sp_info.get("area")) and base_dir and os.path.exists(base_dir):
        try:
            import arcpy
            gdbs = [os.path.join(base_dir, d) for d in os.listdir(base_dir) if d.endswith(".gdb")]
            if gdbs:
                gdb = gdbs[0]
                arcpy.env.workspace = gdb
                fcs = arcpy.ListFeatureClasses() or []
                sname_lower = str(study_area_name).lower()
                bnd_fcs = [fc for fc in fcs if fc.lower() in (sname_lower, str(sname_lower) + "_boundary", "boundary", "study_area", "mask")]
                if not bnd_fcs:
                    for fc in fcs:
                        try:
                            if arcpy.Describe(fc).shapeType == "Polygon":
                                bnd_fcs.append(fc)
                                break
                        except Exception:
                            pass
                if bnd_fcs:
                    b_fc = os.path.join(gdb, bnd_fcs[0])
                    total_sqkm = 0.0
                    total_perim_km = 0.0
                    with arcpy.da.SearchCursor(b_fc, ["SHAPE@"]) as cur:
                        for row in cur:
                            g = row[0]
                            if g:
                                try:
                                    total_sqkm += g.getArea("GEODESIC", "SQUAREKILOMETERS")
                                    total_perim_km += g.getLength("GEODESIC", "KILOMETERS")
                                except Exception:
                                    total_sqkm += (g.area / 1e6)
                                    total_perim_km += (g.length / 1e3)
                    if total_sqkm > 0:
                        sp_info["area"] = "{:,.2f} كم² ({:,.2f} مليون هكتار)".format(total_sqkm, total_sqkm / 10000.0) if total_sqkm >= 10000 else "{:,.2f} كم² ({:,.2f} هكتار)".format(total_sqkm, total_sqkm * 100)
                    if total_perim_km > 0:
                        sp_info["perimeter"] = "{:,.2f} كم".format(total_perim_km)
                    
                    desc = arcpy.Describe(b_fc)
                    ext = desc.extent
                    sr = desc.spatialReference
                    sr_wgs = arcpy.SpatialReference(4326)
                    p_min = arcpy.PointGeometry(arcpy.Point(ext.XMin, ext.YMin), sr).projectAs(sr_wgs)
                    p_max = arcpy.PointGeometry(arcpy.Point(ext.XMax, ext.YMax), sr).projectAs(sr_wgs)
                    sp_info["extent"] = u"{:.4f}°N إلى {:.4f}°N | {:.4f}°E إلى {:.4f}°E".format(p_min.firstPoint.Y, p_max.firstPoint.Y, p_min.firstPoint.X, p_max.firstPoint.X)
                
                pt_fcs = [fc for fc in fcs if "point" in fc.lower() or "grid" in fc.lower() or "sample" in fc.lower()]
                if pt_fcs:
                    p_fc = os.path.join(gdb, pt_fcs[0])
                    cnt = int(arcpy.management.GetCount(p_fc)[0])
                    sp_info["points"] = u"{:,} محطة رصد مناخية".format(cnt)
        except Exception:
            pass

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # ==============================================================================
    # SHEET 1: 01_عن_الدليل_About_Guide
    # ==============================================================================
    ws1 = wb.create_sheet(title="01_عن_الدليل_About_Guide")
    ws1.views.sheetView[0].showGridLines = True

    # Title Banner
    ws1.merge_cells("A1:G1")
    study_display = sp_info.get("display_name") or study_area_name
    tcell = ws1.cell(1, 1, u"أطلس المناخ الرقمي لـ {} — دليل التصنيف الكارتوجرافي الاسترشادي\nDigital Climate Atlas of {} — Advisory Cartographic Classification Guide".format(study_display, study_area_name))
    tcell.font = Font(name="Segoe UI", size=13, bold=True, color="FFFFFF")
    tcell.fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    tcell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws1.row_dimensions[1].height = 45

    # Section 1: Spatial & Boundary Profile
    ws1.merge_cells("A3:G3")
    s1 = ws1.cell(3, 1, "1. بطاقة الهوية المكانية لمنطقة الدراسة (Spatial & Boundary Profile)")
    s1.font = Font(name="Segoe UI", size=11, bold=True, color="1F4E79")
    s1.fill = SECTION_FILL
    s1.alignment = Alignment(horizontal="right", vertical="center")
    ws1.row_dimensions[3].height = 24

    def_area = "محسوبة هندسياً من مضلع قناع منطقة الدراسة"
    def_perim = "محسوبة هندسياً من مضلع قناع منطقة الدراسة"
    def_extent = "مستخرج من نطاق حدود منطقة الدراسة WGS84"
    def_points = "محسوبة من شبكة نقاط الرصد المعتمدة"

    spatial_meta = [
        ("اسم منطقة الدراسة (Study Area Name)", sp_info.get("name", study_area_name), "النطاق الإقليمي المعتمد في التحميل والمعالجة"),
        ("المساحة الكلية لمنطقة الدراسة (Total Area)", sp_info.get("area", def_area), "محسوبة هندسياً بدقة من مضلع قناع منطقة الدراسة"),
        ("المحيط الكلي للحدود (Total Perimeter)", sp_info.get("perimeter", def_perim), "شاملاً الحدود المحيطة لمنطقة الدراسة"),
        ("الامتداد الجغرافي الفلكي (Coordinates Extent)", sp_info.get("extent", def_extent), "النطاق المرجعي العالمي WGS 1984 (EPSG:4326)"),
        ("عدد محطات / نقاط شبكة الرصد (Sampling Grid)", sp_info.get("points", def_points), "موزعة بانتظام جغرافي على كامل نطاق منطقة الدراسة"),
    ]

    for idx, (prop, val, note) in enumerate(spatial_meta, 4):
        ws1.cell(idx, 1, idx - 3).alignment = Alignment(horizontal="center", vertical="center")
        ws1.cell(idx, 1).font = DATA_FONT_BOLD
        ws1.cell(idx, 2, prop).alignment = Alignment(horizontal="right", vertical="center")
        ws1.cell(idx, 2).font = DATA_FONT_BOLD
        ws1.merge_cells(start_row=idx, start_column=3, end_row=idx, end_column=4)
        c3 = ws1.cell(idx, 3, val)
        c3.alignment = Alignment(horizontal="center", vertical="center")
        c3.font = CODE_FONT_BOLD
        ws1.merge_cells(start_row=idx, start_column=5, end_row=idx, end_column=7)
        c5 = ws1.cell(idx, 5, note)
        c5.alignment = Alignment(horizontal="right", vertical="center")
        c5.font = DATA_FONT
        for c in range(1, 8):
            ws1.cell(idx, c).border = THIN_BORDER
            ws1.cell(idx, c).fill = ROW_EVEN_FILL if idx % 2 == 0 else ROW_WHITE_FILL
        ws1.row_dimensions[idx].height = 22

    # Section 2: Technical & Raster Processing Profile
    r_start = 10
    ws1.merge_cells("A{0}:G{0}".format(r_start))
    s2 = ws1.cell(r_start, 1, "2. مواصفات المعالجة الفنية والكارتوجرافية (Technical Processing Parameters)")
    s2.font = Font(name="Segoe UI", size=11, bold=True, color="1F4E79")
    s2.fill = SECTION_FILL
    s2.alignment = Alignment(horizontal="right", vertical="center")
    ws1.row_dimensions[r_start].height = 24

    tc_info = tech_info or {}
    tech_meta = [
        ("دقة خلية الراستر الكارتوجرافية (Raster Cell Size)", tc_info.get("cell_size", "500 متر (500 Meters)"), "مخرجات شبكية منتظمة ملساء خالية من حواف التكسير البصري"),
        ("تباعد شبكة متجهات الرياح (Wind Vector Spacing)", tc_info.get("wind_cell", "25,000 متر (25 كم)"), "شبكة أسهم متجهات اتجاه وسرعة الرياح السطحية"),
        ("فاصل خطوط تساوي الضغط (Isobar Step)", tc_info.get("isobar_step", "4.0 mbar (Global Standard)"), "فاصل الآيزوبار المعتمد أرصادياً عالمياً"),
        ("طريقة الاستيفاء المكاني (Interpolation Algorithm)", tc_info.get("interp", "IDW (Inverse Distance Weighted)"), "خوارزمية مقلوب مربع المسافة الموزونة"),
        ("تنسيق ونمط ملفات الراستر (Raster Format)", tc_info.get("format", "GeoTIFF 32-bit Float (LZW Compressed)"), "راسترات عائمة أصلية مقصوصة تماماً على حدود منطقة الدراسة"),
        ("فترة الرصد المناخي القياسية (Climate Period)", tc_info.get("period", "1996 – 2025 (30 عاماً)"), "تطابق المعدل المناخي القياسي العالمي (Climate Normal)"),
        ("مزود ومصدر البيانات المناخية (Data Source)", tc_info.get("source", "NASA POWER API (CERES & MERRA-2)"), "نماذج إعادة التحليل والرصد الفضائي الدقيق لناسا"),
    ]

    for idx, (prop, val, note) in enumerate(tech_meta, r_start + 1):
        ws1.cell(idx, 1, idx - r_start).alignment = Alignment(horizontal="center", vertical="center")
        ws1.cell(idx, 1).font = DATA_FONT_BOLD
        ws1.cell(idx, 2, prop).alignment = Alignment(horizontal="right", vertical="center")
        ws1.cell(idx, 2).font = DATA_FONT_BOLD
        ws1.merge_cells(start_row=idx, start_column=3, end_row=idx, end_column=4)
        c3 = ws1.cell(idx, 3, val)
        c3.alignment = Alignment(horizontal="center", vertical="center")
        c3.font = CODE_FONT_BOLD
        ws1.merge_cells(start_row=idx, start_column=5, end_row=idx, end_column=7)
        c5 = ws1.cell(idx, 5, note)
        c5.alignment = Alignment(horizontal="right", vertical="center")
        c5.font = DATA_FONT
        for c in range(1, 8):
            ws1.cell(idx, c).border = THIN_BORDER
            ws1.cell(idx, c).fill = ROW_EVEN_FILL if idx % 2 == 0 else ROW_WHITE_FILL
        ws1.row_dimensions[idx].height = 22

    # Section 3: Advisory Guide Philosophy
    r_rec = 19
    ws1.merge_cells("A{0}:G{0}".format(r_rec))
    s3 = ws1.cell(r_rec, 1, "3. فلسفة الدليل الاسترشادي الكارتوجرافي ومستويات الفئات (Advisory Classification Tiers)")
    s3.font = Font(name="Segoe UI", size=11, bold=True, color="1F4E79")
    s3.fill = SECTION_FILL
    s3.alignment = Alignment(horizontal="right", vertical="center")
    ws1.row_dimensions[r_rec].height = 24

    variance_guide_text = (
        "هذا الدليل هو مرجع كارتوجرافي استرشادي (Advisory Guidelines) يُتيح للباحث والمصمم الكارتوجرافي مرونة كاملة في اختيار مستوى الفئات "
        "الأنسب لطبيعة دراسته ومنطقة عمله، دون فرض قوالب جامدة أو إلزامية:\n"
        "• تم تخصيص صفحة مستقلة لكل عنصر مناخي في أسفل المصنف (من 01_Temperature إلى 19_Trends_And_Anomalies).\n"
        "• داخل كل عنصر، تُعرض المؤشرات مقسمة إلى 4 مستويات متتالية (3 فئات، ثم فاصل، 5 فئات، ثم فاصل، 7 فئات، ثم فاصل، 9 فئات).\n"
        "• تم تطبيق الدقة العشرية الذكية الحساسة (Adaptive Decimals) لمنع التقريب العشوائي للأعداد الصحيحة الذي قد يُصفّر أو يشوه المؤشرات الحساسة (مثل مؤشرات الجفاف والاتجاهات العشرية والضغط الجوي).\n"
        "• تم الحفاظ التام على المعايير العالمية الرسمية (Global Standards) للمؤشرات القياسية (مثل معيار WHO للأشعة فوق البنفسجية UV، ومعيار UNEP لأطلس التصحر)."
    )

    rec_rows = [
        ("3 فئات (Executive Overview)", "3", "المعيار الأساسي للتقارير والملخصات التنفيذية ★",
         "البيانات ذات المدى الرقمي الضيق والتقارير التنفيذية الموجزة لمتخذي القرار وصناع السياسات.",
         "وضوح كارتوجرافي فائق، قراءة بصرية فورية وسهلة، ومنع تشتت القارئ غير المتخصص.",
         "تناسب العرض العام ولكنها تدمج بعض التدرجات والتفاصيل المكانية الدقيقة.",
         "متاح في جدول المستوى الأول داخل صفحة كل عنصر"),

        ("5 فئات (Standard & Balanced)", "5", "المعيار الكارتوجرافي الذهبي العام (الأكثر توازناً) ★",
         "الخيار القياسي المعتمد لأكثر من 90% من خرائط أطلس المناخ للباحثين والمخططين والجامعات.",
         "التوازن الأمثل بين إبراز التباين المكاني وسهولة تمييز درجات الألوان وتوافقها مع المعايير الدولية.",
         "المعيار الأكثر شمولاً واعتماداً لمعظم الظواهر المناخية.",
         "متاح في جدول المستوى الثاني داخل صفحة كل عنصر"),

        ("7 فئات (Academic Detailed)", "7", "موصى به في حالات التباين الرقمي والتضاريسي الكبير",
         "المتغيرات ذات التباين المكاني الواسع، ومناطق التضاريس الجبلية والساحلية، والأبحاث الأكاديمية المتخصصة.",
         "إبراز التدرجات الدقيقة والتغيرات الانتقالية الحساسة عبر نطاق منطقة الدراسة.",
         "تتطلب تمايزاً حقيقياً في البيانات لتجنب التقارب اللوني المربك في الطباعة.",
         "متاح في جدول المستوى الثالث داخل صفحة كل عنصر"),

        ("9 فئات (High-Resolution)", "9", "خاص بالتباين الشديد جداً والمدى الرقمي فائق الاتساع",
         "المتغيرات ذات المدى الرقمي الفائق (مثل الإشعاع الشمسي السنوي، أو التباين الحاد في الأمطار والرياح).",
         "تغطية طيفية فائقة الدقة تُظهر أدق الفروق المكانية للبيانات المستمرة عبر مساحات واسعة.",
         "تتطلب شاشات عالية الدقة أو طباعة كارتوجرافية جدارية كبيرة لتمييز درجات الألوان التسع.",
         "متاح في جدول المستوى الرابع داخل صفحة كل عنصر"),
    ]

    ws1.merge_cells("A20:G20")
    rec_desc = ws1.cell(20, 1, variance_guide_text)
    rec_desc.font = Font(name="Segoe UI", size=9.5, italic=True)
    rec_desc.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
    ws1.row_dimensions[20].height = 65

    rec_headers = [
        (1, "خيار التصنيف\nTier", 16),
        (2, "عدد الفئات\nClasses", 12),
        (3, "مستوى التوصية\nRecommendation", 22),
        (4, "أفضل حالات الاستخدام والتطبيق\nBest Use Case", 42),
        (5, "المزايا والخصائص الكارتوجرافية\nCartographic Advantages", 38),
        (6, "المحددات والملاحظات الفنية\nLimitations & Warnings", 38),
        (7, "الموقع في الملف\nLocation in File", 30),
    ]

    for col_idx, h_text, width in rec_headers:
        cell = ws1.cell(22, col_idx, h_text)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN
        cell.border = THIN_BORDER
        col_letter = get_column_letter(col_idx)
        ws1.column_dimensions[col_letter].width = max(ws1.column_dimensions[col_letter].width or 0, width)
    ws1.row_dimensions[22].height = 28

    for r_idx, r_vals in enumerate(rec_rows, 23):
        is_rec = "المعيار الذهبي" in r_vals[2]
        bg = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid") if is_rec else (ROW_EVEN_FILL if r_idx % 2 == 0 else ROW_WHITE_FILL)
        for c_idx, val in enumerate(r_vals, 1):
            cell = ws1.cell(r_idx, c_idx, val)
            cell.border = THIN_BORDER
            cell.fill = bg
            if c_idx in (1, 2, 7):
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.font = DATA_FONT_BOLD if c_idx == 1 else (CODE_FONT_BOLD if c_idx == 7 else DATA_FONT)
            elif c_idx == 3:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.font = Font(name="Segoe UI", size=9, bold=True, color="276A3C" if is_rec else "1F4E79")
            else:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.font = DATA_FONT
        ws1.row_dimensions[r_idx].height = 25

    # ==============================================================================
    # SHEET 2: 02_حقول_الإدارة_Admin_Fields
    # ==============================================================================
    ws2 = wb.create_sheet(title="02_حقول_الإدارة_Admin_Fields")
    ws2.views.sheetView[0].showGridLines = True

    admin_headers = [
        (1, "م\nNo.", 6),
        (2, "اسم الحقل المختصر (Shapefile)\nShort Field Name", 22),
        (3, "اسم الحقل الكامل (Geodatabase)\nFull Field Name (EN)", 26),
        (4, "نوع البيانات\nData Type", 14),
        (5, "وحدة القياس\nUnit", 14),
        (6, "الشرح والتوصيف العلمي باللغة العربية\nArabic Scientific Description", 55),
        (7, "الوصف الفني باللغة الإنجليزية\nEnglish Technical Description", 50),
    ]

    for col_idx, h_text, width in admin_headers:
        cell = ws2.cell(1, col_idx, h_text)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN
        cell.border = THIN_BORDER
        col_letter = get_column_letter(col_idx)
        ws2.column_dimensions[col_letter].width = width
    ws2.row_dimensions[1].height = 30

    admin_meta_expanded = [
        (1, "OBJECTID", "OBJECTID", "Integer (OID)", "-", "المعرف الرقمي الفريد التلقائي لطبقة المعالم داخل قاعدة البيانات الجيوداتابيز.", "Unique primary key identifier generated by File Geodatabase."),
        (2, "Source_ID", "Source_ID", "Integer", "-", "المعرف المرجعي الأصلي الثابت لنقطة شبكة الرصد المناخي.", "Stable reference ID of the meteorological sampling grid point."),
        (3, "Point_Lat", "Point_Lat", "Double", "درجة (°)", "دائرة العرض الجغرافية للنقطة بنظام الإحداثيات العالمي WGS 1984.", "Geographic Latitude in decimal degrees (WGS 1984 / EPSG:4326)."),
        (4, "Point_Lon", "Point_Lon", "Double", "درجة (°)", "خط الطول الجغرافي للنقطة بنظام الإحداثيات العالمي WGS 1984.", "Geographic Longitude in decimal degrees (WGS 1984 / EPSG:4326)."),
        (5, "Data_Start", "Data_Start", "Text (Date)", "سنة / شهر", "تاريخ وسنة بداية سلسلة الرصد المناخي المسترجعة من سيرفر ناسا.", "Start date/year of the retrieved NASA POWER climate observation series."),
        (6, "Data_End", "Data_End", "Text (Date)", "سنة / شهر", "تاريخ وسنة نهاية سلسلة الرصد المناخي المسترجعة من سيرفر ناسا.", "End date/year of the retrieved NASA POWER climate observation series."),
        (7, "Temporal", "Temporal", "Text", "فترة", "النمط والمدى الزمني المستخدم في المعالجة (شهري / يومي / معدلات قياسية).", "Temporal processing resolution (Daily / Monthly / Climatological Normals)."),
        (8, "Intrp_Meth", "Interp_Meth", "Text", "خوارزمية", "خوارزمية الاستيفاء والاشتقاق المكاني المطبقة لإنشاء أسطح الراستر.", "Spatial interpolation method applied (IDW / Kriging / Spline)."),
        (9, "Cell_Size", "Cell_Size", "Double", "متر / درجة", "حجم ودقة خلية الراستر الكارتوجرافية المنتجة.", "Output cartographic raster pixel resolution in linear meters or degrees."),
        (10, "Wind_Cell", "Wind_Cell", "Double", "متر / درجة", "تباعد شبكة متجهات وأسهم الرياح الكارتوجرافية.", "Cartographic fishnet spacing for wind vector arrows and direction points."),
        (11, "Measurement_Unit", "Measurement_Unit", "Text (25)", "رمز الوحدة", "وحدة القياس المعتمدة للمتغير في الطبقة (°C، mm، hPa، m/s، إلخ).", "Standard measurement unit assigned to the climate variable in the layer."),
        (12, "Status", "Status", "Text", "حالة", "حالة نجاح استرجاع ومعالجة بيانات النقطة من السيرفر (OK / FAILED).", "Data query status for the station point (OK or FAILED)."),
        (13, "Error_Msg", "Error_Msg", "Text", "رسالة", "رسالة الخطأ التفصيلية في حال تعذر الاتصال بالنقطة أو وجود قيم غير صالحة.", "Detailed error description if query timed out or returned invalid data."),
    ]

    for row_idx, rdata in enumerate(admin_meta_expanded, 2):
        for col_idx, val in enumerate(rdata, 1):
            cell = ws2.cell(row_idx, col_idx, val)
            cell.border = THIN_BORDER
            cell.fill = ROW_EVEN_FILL if row_idx % 2 == 0 else ROW_WHITE_FILL
            if col_idx in (1, 4, 5):
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.font = DATA_FONT_BOLD if col_idx == 1 else DATA_FONT
            elif col_idx in (2, 3):
                cell.alignment = Alignment(horizontal="left", vertical="center")
                cell.font = CODE_FONT_BOLD if col_idx == 2 else CODE_FONT
            elif col_idx == 6:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.font = DATA_FONT
            elif col_idx == 7:
                cell.alignment = Alignment(horizontal="left", vertical="center")
                cell.font = DATA_FONT
        ws2.row_dimensions[row_idx].height = 24

    # ==============================================================================
    # SHEET 3: 03_المعايير_القياسية_Standards
    # ==============================================================================
    ws3 = wb.create_sheet(title="03_المعايير_القياسية_Standards")
    ws3.views.sheetView[0].showGridLines = True

    standards_headers = [
        (1, "الرقم\nNo.", 6),
        (2, "اسم العنصر في الأداة\nTool Parameter & Layer", 25),
        (3, "الاسم باللغة العربية\nModule Arabic Name", 28),
        (4, "مجلد الراستر (1:1)\nOutput Raster Folder", 26),
        (5, "طبقة المعالم (1:1)\nFeature Class Name", 26),
        (6, "المؤشرات\nIndicators", 12),
        (7, "الوحدة\nMain Unit", 14),
        (8, "المنهجية والمعادلة العلمية (بالعربية)\nArabic Scientific Methodology & Formula", 55),
        (9, "المعايير والمراجع الدولية (بالإنجليزية)\nScientific Method & Standards (EN)", 45),
    ]

    for col_idx, h_text, width in standards_headers:
        cell = ws3.cell(1, col_idx, h_text)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN
        cell.border = THIN_BORDER
        col_letter = get_column_letter(col_idx)
        ws3.column_dimensions[col_letter].width = width
    ws3.row_dimensions[1].height = 32

    has_project_rasters = any(f['full_f'] in raster_stats for f in climate_fields)

    for row_idx, mod_item in enumerate(MODULES_METADATA, 2):
        # Count actual indicators belonging to this module
        pfx = mod_item["pfx"]
        code = mod_item["code"]
        if has_project_rasters:
            m_count = len([f for f in climate_fields if mod_field_matches(f['module'], pfx, code) and f['full_f'] in raster_stats])
        else:
            m_count = len([f for f in climate_fields if mod_field_matches(f['module'], pfx, code)])
        rdata = (
            pfx,
            mod_item["code"],
            mod_item["name_ar"],
            mod_item["folder"],
            mod_item["fc"],
            m_count if m_count > 0 else 1,
            mod_item["unit"],
            mod_item["desc_ar"],
            mod_item["method_en"]
        )
        for col_idx, val in enumerate(rdata, 1):
            cell = ws3.cell(row_idx, col_idx, val)
            cell.border = THIN_BORDER
            cell.fill = ROW_EVEN_FILL if row_idx % 2 == 0 else ROW_WHITE_FILL
            if col_idx in (1, 6, 7):
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.font = DATA_FONT_BOLD if col_idx == 1 else DATA_FONT
            elif col_idx in (2, 4, 5, 9):
                cell.alignment = Alignment(horizontal="left", vertical="center")
                cell.font = CODE_FONT_BOLD if col_idx in (2, 4, 5) else DATA_FONT
            elif col_idx in (3, 8):
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.font = DATA_FONT_BOLD if col_idx == 3 else DATA_FONT
        ws3.row_dimensions[row_idx].height = 28

    # Total Summary row in Standards
    total_r = len(MODULES_METADATA) + 2
    ws3.cell(total_r, 3, "الإجمالي العام للمؤشرات المناخية الخرائطية:").alignment = Alignment(horizontal="right", vertical="center")
    ws3.cell(total_r, 3).font = Font(name="Segoe UI", size=10, bold=True, color="1F4E79")
    if has_project_rasters:
        tot_indicators = sum(len([f for f in climate_fields if mod_field_matches(f['module'], m['pfx'], m['code']) and f['full_f'] in raster_stats]) for m in MODULES_METADATA)
    else:
        tot_indicators = len(climate_fields)
    ws3.cell(total_r, 6, tot_indicators).alignment = Alignment(horizontal="center", vertical="center")
    ws3.cell(total_r, 6).font = Font(name="Consolas", size=10, bold=True, color="1F4E79")
    for c in range(1, 10):
        ws3.cell(total_r, c).border = THIN_BORDER
        ws3.cell(total_r, c).fill = SECTION_FILL
    ws3.row_dimensions[total_r].height = 25

    # ==============================================================================
    # SHEETS 4..22: THE 19 DEDICATED CLIMATE MODULE SHEETS
    # ==============================================================================
    # Level tier definitions
    TIERS = [
        {
            "n": 3,
            "banner_text": u"🔷 المستوى الأول: تصنيف 3 فئات استرشادية (نطاق عام ملخص للإدارة وصناع القرار — Level 1: 3 Classes Overview)",
            "banner_fill": PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid"),
            "header_fill": PatternFill(start_color="2C5E8A", end_color="2C5E8A", fill_type="solid")
        },
        {
            "n": 5,
            "banner_text": u"🟢 المستوى الثاني: تصنيف 5 فئات استرشادية (المعيار القياسي الذهبي المتوازن للأطلس المناخي — Level 2: 5 Classes Recommended)",
            "banner_fill": PatternFill(start_color="196F3D", end_color="196F3D", fill_type="solid"),
            "header_fill": PatternFill(start_color="239B56", end_color="239B56", fill_type="solid")
        },
        {
            "n": 7,
            "banner_text": u"🟠 المستوى الثالث: تصنيف 7 فئات استرشادية (المعيار الأكاديمي والتفصيلي للدراسات والبحوث المتخصصة — Level 3: 7 Classes Detailed)",
            "banner_fill": PatternFill(start_color="B9770E", end_color="B9770E", fill_type="solid"),
            "header_fill": PatternFill(start_color="D68910", end_color="D68910", fill_type="solid")
        },
        {
            "n": 9,
            "banner_text": u"🟣 المستوى الرابع: تصنيف 9 فئات استرشادية (المعيار الموسع للمناطق ذات التباين التضاريسي الواسع — Level 4: 9 Classes Wide Range)",
            "banner_fill": PatternFill(start_color="6C3483", end_color="6C3483", fill_type="solid"),
            "header_fill": PatternFill(start_color="884EA0", end_color="884EA0", fill_type="solid")
        }
    ]

    for mod_info in MODULES_METADATA:
        m_code = mod_info["code"]
        m_pfx = mod_info["pfx"]
        m_name_ar = mod_info["name_ar"]
        m_name_en = mod_info["name_en"]
        m_unit = mod_info["unit"]
        m_folder = mod_info["folder"]

        # Filter fields belonging to this module
        if has_project_rasters:
            mod_fields = [f for f in climate_fields if mod_field_matches(f['module'], m_pfx, m_code) and f['full_f'] in raster_stats]
        else:
            mod_fields = [f for f in climate_fields if mod_field_matches(f['module'], m_pfx, m_code)]
        if not mod_fields:
            continue

        ws_mod = wb.create_sheet(title=m_code)
        ws_mod.views.sheetView[0].showGridLines = True

        # Module Header Card (Rows 1-3)
        ws_mod.merge_cells("A1:M1")
        t_card = ws_mod.cell(1, 1, u"أطلس المناخ الرقمي لـ {} — دليل تصنيف مؤشرات {} ({})\nDigital Climate Atlas of {} — {} Module Classification Guide".format(study_display, m_name_ar, m_code, study_area_name, m_name_en))
        t_card.font = Font(name="Segoe UI", size=13, bold=True, color="FFFFFF")
        t_card.fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        t_card.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws_mod.row_dimensions[1].height = 42

        ws_mod.merge_cells("A2:M2")
        sub_card = ws_mod.cell(2, 1, u"المجال المناخي: {} ({}) | الوحدة الأساسية: {} | عدد المؤشرات: {} مؤشراً | المستويات الاسترشادية: 3 و 5 و 7 و 9 فئات".format(m_name_ar, m_name_en, m_unit, len(mod_fields)))
        sub_card.font = Font(name="Segoe UI", size=9.5, bold=True, color="1F4E79")
        sub_card.fill = PatternFill(start_color="EAECEE", end_color="EAECEE", fill_type="solid")
        sub_card.alignment = Alignment(horizontal="center", vertical="center")
        ws_mod.row_dimensions[2].height = 22

        ws_mod.row_dimensions[3].height = 10
        curr_row = 4

        # Iterate over the 4 Tiers (3, 5, 7, 9 classes)
        for tier in TIERS:
            n_classes = tier["n"]
            banner_text = tier["banner_text"]
            b_fill = tier["banner_fill"]
            h_fill = tier["header_fill"]

            max_col = 10 + n_classes

            # 1. Section Banner Row
            ws_mod.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=max_col)
            b_cell = ws_mod.cell(curr_row, 1, banner_text)
            b_cell.font = Font(name="Segoe UI", size=10.5, bold=True, color="FFFFFF")
            b_cell.fill = b_fill
            b_cell.alignment = Alignment(horizontal="center", vertical="center")
            for c in range(1, max_col + 1):
                ws_mod.cell(curr_row, c).border = THIN_BORDER
                ws_mod.cell(curr_row, c).fill = b_fill
            ws_mod.row_dimensions[curr_row].height = 26
            curr_row += 1

            # 2. Table Header Row
            headers = [
                (1, "م\nNo.", 6),
                (2, "اسم الحقل في قاعدة البيانات\nField Name (EN)", 26),
                (3, "اسم الخريطة المقترح للأطلس (بالعربية)\nSuggested Atlas Map Title (AR)", 36),
                (4, "التوصيف العلمي باللغة العربية\nArabic Scientific Description", 48),
                (5, "وحدة القياس\nUnit", 14),
                (6, "طريقة التصنيف\nClassification Method", 28),
                (7, "طول الخطوة\nStep", 16),
                (8, "أدنى قيمة بالمنطقة\nMin", 15),
                (9, "أقصى قيمة بالمنطقة\nMax", 15),
                (10, "متوسط المنطقة\nMean", 15),
            ]

            for ci in range(1, n_classes + 1):
                if ci == 1:
                    lbl = u"الفئة 1 (الأعلى)\nClass 1 (Highest)"
                elif ci == n_classes:
                    lbl = u"الفئة {} (الأدنى)\nClass {} (Lowest)".format(ci, ci)
                else:
                    lbl = u"الفئة {}\nClass {}".format(ci, ci)
                headers.append((10 + ci, lbl, 22))

            for col_idx, h_text, width in headers:
                cell = ws_mod.cell(curr_row, col_idx, h_text)
                cell.font = HEADER_FONT
                cell.fill = h_fill
                cell.alignment = HEADER_ALIGN
                cell.border = THIN_BORDER
                col_letter = get_column_letter(col_idx)
                cur_w = ws_mod.column_dimensions[col_letter].width or 0
                if width > cur_w:
                    ws_mod.column_dimensions[col_letter].width = width
            ws_mod.row_dimensions[curr_row].height = 32
            curr_row += 1

            # 3. Data Rows
            for map_idx, fentry in enumerate(mod_fields, 1):
                full_f = fentry['full_f']
                unit = fentry['unit']
                title_ar = fentry['title_ar']
                desc_ar = fentry['desc_ar']

                # Lookup actual raster stats
                st = raster_stats.get(full_f, {'min': 0.0, 'max': 1.0, 'mean': 0.5, 'tif': full_f + '.tif'})
                vmin = float(st['min'])
                vmax = float(st['max'])
                vmean = float(st['mean'])

                # Physical floor for strictly non-negative meteorological variables
                if full_f.startswith(("UNEP_", "DM_", "R_", "ET_", "PET_", "W_Spd", "Sol_", "UV_", "Cld_", "RH_", "Dry_Months")):
                    if vmin < 0.0:
                        vmin = 0.0
                    if vmax < 0.0:
                        vmax = 0.0
                    if vmean < 0.0:
                        vmean = 0.0

                # Compute classes with adaptive precision & global standards
                step_str, method_str, labels, prec = compute_smart_classes(full_f, vmin, vmax, n_classes, unit)

                # Clean zero rounding to avoid -0.000
                if round(vmin, prec) == 0.0:
                    vmin = 0.0
                if round(vmax, prec) == 0.0:
                    vmax = 0.0
                if round(vmean, prec) == 0.0:
                    vmean = 0.0

                u_sym = get_clean_unit_symbol(unit, full_f)
                u_str = (" " + u_sym) if u_sym and not u_sym.startswith("°") and u_sym != "%" else (u_sym or "")

                if prec == 0:
                    min_str = "{}{}".format(int(round(vmin)), u_str)
                    max_str = "{}{}".format(int(round(vmax)), u_str)
                    mean_str = "{}{}".format(int(round(vmean)), u_str)
                else:
                    min_str = "{:.{}f}{}".format(vmin, prec, u_str)
                    max_str = "{:.{}f}{}".format(vmax, prec, u_str)
                    mean_str = "{:.{}f}{}".format(vmean, prec, u_str)

                bg_fill = ROW_EVEN_FILL if map_idx % 2 == 0 else ROW_WHITE_FILL

                row_cells = [
                    (1, map_idx, Alignment(horizontal="center", vertical="center"), DATA_FONT_BOLD),
                    (2, full_f, Alignment(horizontal="left", vertical="center"), CODE_FONT_BOLD),
                    (3, title_ar, Alignment(horizontal="right", vertical="center"), DATA_FONT_BOLD),
                    (4, desc_ar, Alignment(horizontal="right", vertical="center"), DATA_FONT),
                    (5, u_sym or unit, Alignment(horizontal="center", vertical="center"), DATA_FONT_BOLD),
                    (6, method_str, Alignment(horizontal="center", vertical="center"), DATA_FONT),
                    (7, step_str, Alignment(horizontal="center", vertical="center"), CODE_FONT_BOLD),
                    (8, min_str, Alignment(horizontal="center", vertical="center"), CODE_FONT),
                    (9, max_str, Alignment(horizontal="center", vertical="center"), CODE_FONT_BOLD),
                    (10, mean_str, Alignment(horizontal="center", vertical="center"), CODE_FONT),
                ]

                for ci in range(1, n_classes + 1):
                    lbl_c = labels[ci - 1] if ci <= len(labels) else ""
                    is_extreme = (ci == 1 or ci == n_classes)
                    fnt_c = DATA_FONT_BOLD if is_extreme else DATA_FONT
                    row_cells.append((10 + ci, lbl_c, Alignment(horizontal="center", vertical="center", wrap_text=True), fnt_c))

                for cidx, val, align, fnt in row_cells:
                    cell = ws_mod.cell(curr_row, cidx, val)
                    cell.alignment = align
                    cell.font = fnt
                    cell.fill = bg_fill
                    cell.border = THIN_BORDER
                ws_mod.row_dimensions[curr_row].height = 24
                curr_row += 1

            # 4. Spacer row between tiers
            ws_mod.row_dimensions[curr_row].height = 16
            curr_row += 1

    # Save primary copy to target_excel (typically 00_Tables_And_Reports\Climate_Atlas_Classification_Guide.xlsx)
    os.makedirs(os.path.dirname(target_excel), exist_ok=True)
    wb.save(target_excel)
    print("Successfully generated Master Climate Atlas Classification Guide: " + str(target_excel))

    # Save root workspace copy directly to base_dir
    root_copy = os.path.join(base_dir, "Climate_Atlas_Classification_Guide.xlsx")
    try:
        shutil.copy2(target_excel, root_copy)
        print("Successfully synced root workspace copy: " + str(root_copy))
    except Exception as ex_rc:
        try:
            wb.save(root_copy)
            print("Successfully saved direct root copy: " + str(root_copy))
        except Exception:
            pass

    return target_excel


if __name__ == "__main__":
    import argparse
    import json
    parser = argparse.ArgumentParser()
    parser.add_argument("--json_args", default=None)
    parser.add_argument("--base_dir", default=None)
    parser.add_argument("--study_area", default="Egypt")
    parser.add_argument("--target_excel", default=None)
    args = parser.parse_args()

    if args.json_args and os.path.exists(args.json_args):
        with open(args.json_args, "r") as jf:
            data = json.load(jf)
        build_master_classification_workbook(
            base_dir=data.get("base_dir"),
            study_area_name=data.get("study_area_name", "Egypt"),
            target_excel=data.get("target_excel"),
            spatial_info=data.get("spatial_info"),
            tech_info=data.get("tech_info"),
            raster_stats=data.get("raster_stats")
        )
    else:
        build_master_classification_workbook(
            base_dir=args.base_dir,
            study_area_name=args.study_area,
            target_excel=args.target_excel
        )
